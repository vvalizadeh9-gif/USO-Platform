"""Excel and PDF of the KPI & Performance page.

Both are built from the payload ``services/kpi.summary`` already returned to
the screen, not from a second query. That is the only way "the export matches
the screen" can be a property rather than a hope: there is one set of numbers,
and the file and the page are two renderings of it.

Libraries: ``openpyxl``, which the platform already uses for every other
export, and ``reportlab`` for the PDF. Both are plain Python packages. No new
Docker service, no browser, no external renderer.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# UEP light-mode palette, as the page uses it.
NAVY = "16202E"
MUTED = "5B6778"
TEAL = "0EA394"
TEAL_DARK = "0A7D72"
BORDER = "E2E6EC"
BETTER_STRONG = "9FDDD2"
BETTER = "E3F5F2"
WORSE = "FCE9E7"
WORSE_STRONG = "F4B9B3"
GREY = "EDF0F4"

#: Difference from the country average, in points, at which a cell goes from
#: the soft shade to the strong one. Same number as the page's legend.
STRONG_AT = 5.0

LENS_LABELS = {
    "rm": "Regional Manager",
    "coordinator": "PSO Coordinator",
    "contractor": "Contractor",
    "region": "CRA Region",
}

_HEATMAP_COLUMNS = (
    ("on_air", "On air"),
    ("dt_done", "DT done"),
    ("ict_approved", "ICT approved"),
    ("ict_rejected", "ICT rejected"),
    ("cra_approved", "CRA approved"),
    ("cra_rejected", "CRA rejected"),
)


def _stamp(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return str(value)


def _pct_text(value) -> str:
    return "—" if value is None else f"{value:.1f}%"


def _delta_text(cell: dict, low_sample: bool) -> str:
    if low_sample:
        return "not compared"
    delta = cell.get("delta")
    if delta is None:
        return "—"
    return f"{delta:+.1f} pts"


def cell_shade(cell: dict, low_sample: bool) -> str | None:
    """The colour a heatmap cell gets, or None for no colour.

    Low-sample provinces are never coloured: a province with three DT-done
    villages is not "12 points better than the country", it is three villages.
    For the rejected columns lower is better, so the comparison is inverted —
    otherwise a province that is refused less often than the rest would be
    painted red for it.
    """
    if low_sample:
        return None
    delta = cell.get("delta")
    if delta is None:
        return None
    if cell.get("lower_is_better"):
        delta = -delta
    if delta >= STRONG_AT:
        return BETTER_STRONG
    if delta >= 0:
        return BETTER
    if delta > -STRONG_AT:
        return WORSE
    return WORSE_STRONG


def filename(payload: dict, extension: str) -> str:
    lens = LENS_LABELS.get(payload["lens"], payload["lens"]).replace(" ", "-")
    key = str(payload.get("key") or "all").replace(" ", "-")
    stamp = datetime.now().strftime("%Y%m%d")
    return f"UEP-KPI-{lens}-{key}-{stamp}.{extension}"


def _header_lines(payload: dict) -> list[tuple[str, str]]:
    return [
        ("Lens", LENS_LABELS.get(payload["lens"], payload["lens"])),
        ("Person", str(payload.get("key") or "—")),
        ("Scope", payload["scope"]["chip"]),
        ("Last CPM import", _stamp(payload.get("last_cpm_import"))),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]


# ----- Excel --------------------------------------------------------------


def build_workbook(payload: dict, contractor_block: dict | None) -> bytes:
    """One sheet per section: Summary, Provinces, and Contractors when shown."""
    workbook = Workbook()
    _summary_sheet(workbook.active, payload)
    _provinces_sheet(workbook.create_sheet("Provinces"), payload)
    if contractor_block is not None:
        _contractors_sheet(workbook.create_sheet("Contractors"), contractor_block)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_header(sheet, payload: dict, title: str) -> int:
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14, color=NAVY)
    row = 3
    for label, value in _header_lines(payload):
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, color=MUTED)
        sheet.cell(row=row, column=2, value=value)
        row += 1
    return row + 1


def _section(sheet, row: int, title: str) -> int:
    sheet.cell(row=row, column=1, value=title).font = Font(bold=True, color=TEAL_DARK)
    return row + 1


def _table(sheet, row: int, headers: list[str], rows: list[list]) -> int:
    thin = Side(style="thin", color=BORDER)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index, value=header)
        cell.font = Font(bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=GREY)
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    row += 1
    for values in rows:
        for index, value in enumerate(values, start=1):
            sheet.cell(row=row, column=index, value=value)
        row += 1
    return row + 1


def _summary_sheet(sheet, payload: dict) -> None:
    sheet.title = "Summary"
    row = _write_header(sheet, payload, "KPI & Performance")

    work_items = payload["work_items"]
    row = _section(sheet, row, "Work items")
    row = _table(
        sheet,
        row,
        ["Stage", "Count", "% of scope", "Remain", "Country %", "Difference (pts)"],
        [
            [
                "Total",
                work_items["total"],
                "100.0%" if work_items["total"] else "—",
                0,
                "",
                "",
            ],
            [
                "On air",
                work_items["on_air"],
                _pct_text(work_items["on_air_pct"]),
                work_items["remain_on_air"],
                _pct_text(work_items["country_on_air_pct"]),
                _difference(work_items["on_air_pct"], work_items["country_on_air_pct"]),
            ],
            [
                "DT done",
                work_items["dt_done"],
                _pct_text(work_items["dt_done_pct"]),
                work_items["remain_dt"],
                _pct_text(work_items["country_dt_done_pct"]),
                _difference(
                    work_items["dt_done_pct"], work_items["country_dt_done_pct"]
                ),
            ],
        ],
    )

    villages = payload["villages"]
    row = _section(sheet, row, "Villages")
    row = _table(
        sheet,
        row,
        ["Stage", "Count", "%", "Base", "Remain", "Country %", "Difference (pts)"],
        [
            ["Total", villages["total"], "100.0%" if villages["total"] else "—",
             "scope", 0, "", ""],
            ["On air", villages["on_air"], _pct_text(villages["on_air_pct"]),
             "scope", villages["remain_on_air"],
             _pct_text(villages["country_on_air_pct"]),
             _difference(villages["on_air_pct"], villages["country_on_air_pct"])],
            ["DT done", villages["dt_done"], _pct_text(villages["dt_done_pct"]),
             "scope", villages["remain_dt"],
             _pct_text(villages["country_dt_done_pct"]),
             _difference(villages["dt_done_pct"], villages["country_dt_done_pct"])],
            ["ICT approved", villages["ict_approved"],
             _pct_text(villages["ict_approved_pct"]), "DT-done villages",
             villages["ict_remained"],
             _pct_text(villages["country_ict_approved_pct"]),
             _difference(villages["ict_approved_pct"],
                         villages["country_ict_approved_pct"])],
            ["CRA approved", villages["cra_approved"],
             _pct_text(villages["cra_approved_pct"]), "DT-done villages",
             villages["cra_remained"],
             _pct_text(villages["country_cra_approved_pct"]),
             _difference(villages["cra_approved_pct"],
                         villages["country_cra_approved_pct"])],
        ],
    )

    country = payload["country"]
    row = _section(sheet, row, "Country average (weighted, all 31 provinces)")
    _table(
        sheet,
        row,
        ["Measure", "Numerator", "Denominator", "Weighted %"],
        [
            ["Villages on air", country["villages_on_air"], country["villages"],
             _pct_text(country["on_air_pct"])],
            ["Villages DT done", country["villages_dt_done"], country["villages"],
             _pct_text(country["dt_done_pct"])],
            ["ICT approved", country["ict_approved"], country["villages_dt_done"],
             _pct_text(country["ict_approved_pct"])],
            ["ICT rejected", country["ict_rejected"], country["villages_dt_done"],
             _pct_text(country["ict_rejected_pct"])],
            ["CRA approved", country["cra_approved"], country["villages_dt_done"],
             _pct_text(country["cra_approved_pct"])],
            ["CRA rejected", country["cra_rejected"], country["villages_dt_done"],
             _pct_text(country["cra_rejected_pct"])],
        ],
    )
    _widen(sheet, [24, 16, 16, 18, 14, 14, 18])


def _difference(value, benchmark) -> str:
    if value is None or benchmark is None:
        return "—"
    return f"{value - benchmark:+.1f}"


def _provinces_sheet(sheet, payload: dict) -> None:
    headers = ["Province", "CRA region", "Villages", "DT-done villages"]
    for _key, label in _HEATMAP_COLUMNS:
        headers += [f"{label} %", f"{label} vs country"]

    rows = []
    shades = []
    for province in payload["provinces"] + [payload["country_row"]]:
        low = province["low_sample"]
        values = [
            province["province"],
            province.get("cra_region") or "—",
            province["villages"],
            province["dt_done_villages"],
        ]
        row_shades = [None, None, None, None]
        for key, _label in _HEATMAP_COLUMNS:
            cell = province[key]
            values += [_pct_text(cell["pct"]), _delta_text(cell, low)]
            shade = GREY if province is payload["country_row"] else cell_shade(cell, low)
            row_shades += [shade, shade]
        rows.append(values)
        shades.append(row_shades)

    start = _table(sheet, 1, headers, rows) - len(rows) - 1
    for offset, row_shades in enumerate(shades):
        for index, shade in enumerate(row_shades, start=1):
            if shade:
                sheet.cell(row=start + offset, column=index).fill = PatternFill(
                    "solid", fgColor=shade
                )
    _widen(sheet, [26, 14, 12, 16] + [13, 18] * len(_HEATMAP_COLUMNS))


def _contractors_sheet(sheet, block: dict) -> None:
    mode = block["mode"].upper()
    sheet["A1"] = f"Contractors — {mode}"
    sheet["A1"].font = Font(bold=True, size=13, color=NAVY)
    rows = [
        [r["contractor"], r["villages"], r["approved"], r["remained"], _pct_text(r["pct"])]
        for r in block["rows"]
    ]
    total = block["total"]
    rows.append(
        [total["contractor"], total["villages"], total["approved"],
         total["remained"], _pct_text(total["pct"])]
    )
    _table(
        sheet,
        3,
        ["Contractor (DT SC)", "Total villages", "Approved", "Remained", "Approved %"],
        rows,
    )
    _widen(sheet, [34, 16, 14, 14, 14])


def _widen(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


# ----- PDF ----------------------------------------------------------------


def build_pdf(payload: dict, contractor_block: dict | None) -> bytes:
    """A4 landscape, the same sections as the workbook, in UEP colours.

    Imported inside the function so that a deployment that has not yet
    installed ``reportlab`` still serves the page and the Excel export, and
    fails only on the one request that needs it, with a message that says so.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover - deployment guard
        raise RuntimeError(
            "The PDF export needs the 'reportlab' package, which is listed in "
            "backend/requirements.txt. Rebuild the backend image."
        ) from exc

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="UEP — KPI & Performance",
    )

    navy = colors.HexColor(f"#{NAVY}")
    muted = colors.HexColor(f"#{MUTED}")
    teal_dark = colors.HexColor(f"#{TEAL_DARK}")

    base = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "UepTitle", parent=base["Title"], textColor=navy, fontSize=18, alignment=0
    )
    section_style = ParagraphStyle(
        "UepSection", parent=base["Heading2"], textColor=teal_dark, fontSize=12
    )
    meta_style = ParagraphStyle(
        "UepMeta", parent=base["Normal"], textColor=muted, fontSize=9, leading=13
    )

    story = [Paragraph("KPI &amp; Performance", title_style), Spacer(1, 4)]
    story.append(
        Paragraph(
            " · ".join(f"{label}: {value}" for label, value in _header_lines(payload)),
            meta_style,
        )
    )
    story.append(Spacer(1, 10))

    def grid(data, widths, highlight_last=False, shades=None):
        table = Table(data, colWidths=widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{GREY}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), navy),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("TEXTCOLOR", (0, 1), (-1, -1), navy),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(f"#{BORDER}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if highlight_last:
            style.append(
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{GREY}"))
            )
        for row_index, column_index, shade in shades or []:
            style.append(
                (
                    "BACKGROUND",
                    (column_index, row_index),
                    (column_index, row_index),
                    colors.HexColor(f"#{shade}"),
                )
            )
        table.setStyle(TableStyle(style))
        return table

    work_items = payload["work_items"]
    villages = payload["villages"]
    story.append(Paragraph("Funnels", section_style))
    story.append(
        grid(
            [
                ["Stage", "Work items", "%", "Remain", "Villages", "%", "Remain"],
                [
                    "Scope",
                    work_items["total"],
                    "100.0%" if work_items["total"] else "—",
                    "—",
                    villages["total"],
                    "100.0%" if villages["total"] else "—",
                    "—",
                ],
                [
                    "On air",
                    work_items["on_air"],
                    _pct_text(work_items["on_air_pct"]),
                    work_items["remain_on_air"],
                    villages["on_air"],
                    _pct_text(villages["on_air_pct"]),
                    villages["remain_on_air"],
                ],
                [
                    "DT done",
                    work_items["dt_done"],
                    _pct_text(work_items["dt_done_pct"]),
                    work_items["remain_dt"],
                    villages["dt_done"],
                    _pct_text(villages["dt_done_pct"]),
                    villages["remain_dt"],
                ],
                [
                    "ICT approved",
                    "—",
                    "—",
                    "—",
                    villages["ict_approved"],
                    f'{_pct_text(villages["ict_approved_pct"])} of DT done',
                    villages["ict_remained"],
                ],
                [
                    "CRA approved",
                    "—",
                    "—",
                    "—",
                    villages["cra_approved"],
                    f'{_pct_text(villages["cra_approved_pct"])} of DT done',
                    villages["cra_remained"],
                ],
            ],
            [70 * mm, 28 * mm, 26 * mm, 24 * mm, 28 * mm, 38 * mm, 24 * mm],
        )
    )
    story.append(Spacer(1, 12))

    story.append(Paragraph("Provinces", section_style))
    headers = ["Province", "CRA region", "Villages", "DT done"]
    for _key, label in _HEATMAP_COLUMNS:
        headers.append(label)
    data = [headers]
    shades: list[tuple[int, int, str]] = []
    all_rows = payload["provinces"] + [payload["country_row"]]
    for row_index, province in enumerate(all_rows, start=1):
        low = province["low_sample"]
        line = [
            province["province"],
            province.get("cra_region") or "—",
            province["villages"],
            province["dt_done_villages"],
        ]
        for column_index, (key, _label) in enumerate(_HEATMAP_COLUMNS, start=4):
            cell = province[key]
            line.append(f'{_pct_text(cell["pct"])}\n{_delta_text(cell, low)}')
            shade = (
                GREY
                if province is payload["country_row"]
                else cell_shade(cell, low)
            )
            if shade:
                shades.append((row_index, column_index, shade))
        data.append(line)

    story.append(
        grid(
            data,
            [46 * mm, 24 * mm, 18 * mm, 18 * mm] + [26 * mm] * len(_HEATMAP_COLUMNS),
            highlight_last=True,
            shades=shades,
        )
    )

    if contractor_block is not None:
        story.append(Spacer(1, 12))
        story.append(
            Paragraph(f'Contractors — {contractor_block["mode"].upper()}', section_style)
        )
        rows = [["Contractor (DT SC)", "Total villages", "Approved", "Remained", "Approved %"]]
        for row in contractor_block["rows"]:
            rows.append(
                [row["contractor"], row["villages"], row["approved"],
                 row["remained"], _pct_text(row["pct"])]
            )
        total = contractor_block["total"]
        rows.append(
            [total["contractor"], total["villages"], total["approved"],
             total["remained"], _pct_text(total["pct"])]
        )
        story.append(
            grid(
                rows,
                [80 * mm, 32 * mm, 28 * mm, 28 * mm, 28 * mm],
                highlight_last=True,
            )
        )

    document.build(story)
    return buffer.getvalue()
