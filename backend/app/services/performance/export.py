"""Roles Performance workbooks, built from the payload the screen receives.

Never from a second query: a number in the file cannot differ from the one on
screen, and the main sheet has one row per row the screen shows (the tests
compare the counts). Written in openpyxl's write-only mode into a spooled
file, so a large one is streamed exactly as the Lifecycle Gaps exports are.
"""
from __future__ import annotations

import re
import tempfile
from collections.abc import Iterable
from datetime import date

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from app.core import jalali
from app.services.gap_export import BuiltWorkbook

_BOLD = Font(bold=True)
_SPOOL_BYTES = 8 * 1024 * 1024

Sheet = tuple[str, list[str], Iterable[list]]


def _write(sheets: list[Sheet]) -> BuiltWorkbook:
    """Write the sheets; ``rows`` is the first (main) sheet's row count."""
    wb = Workbook(write_only=True)
    main_rows = None
    for title, header, rows in sheets:
        ws = wb.create_sheet(title)
        ws.freeze_panes = "A2"
        ws.append([_bold(ws, h) for h in header])
        count = 0
        for row in rows:
            ws.append(row)
            count += 1
        if main_rows is None:
            main_rows = count
    target = tempfile.SpooledTemporaryFile(max_size=_SPOOL_BYTES)
    wb.save(target)
    target.seek(0)
    return BuiltWorkbook(file=target, rows=main_rows or 0)


def _bold(ws, value) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = _BOLD
    return cell


def _about(lines: list[tuple[str, object]]) -> Sheet:
    exported = jalali.format_shamsi(jalali.tehran_today())
    return ("About", ["Field", "Value"], [[k, v] for k, v in [*lines, ("Exported", exported)]])


def filename(tab: str, label: str, *, on: date | None = None) -> str:
    """``uep-roles-month-whole-country-1405-07-10.xlsx``: ASCII only."""
    slug = re.sub(r"[^a-z0-9]+", "-", label.lower()).strip("-") or "scope"
    day = jalali.format_shamsi(on or jalali.tehran_today()).replace("/", "-")
    return f"uep-roles-{tab}-{slug}-{day}.xlsx"


# ----- One builder per tab --------------------------------------------------


_OWNER_KINDS = {"contractor": "Contractor", "rm": "Regional manager",
                "coordinator": "Coordinator"}


def month(payload: dict) -> BuiltWorkbook:
    by_owner = payload["by"] not in ("all", "owners")
    figures = ["This month", "Last month (same day)", "Change"]
    if payload["by"] == "owners":
        header = ["Section", "Measure", "Owner kind", "Owner", *figures]
    elif by_owner:
        header = ["Section", "Measure", "Owner", *figures]
    else:
        header = ["Section", "Measure", "Unit", *figures,
                  "Villages this month", "Villages last month", "Villages change"]

    def rows():
        for section in payload["sections"]:
            for row in section["rows"]:
                if row.get("owners_by"):
                    for kind, owners in row["owners_by"].items():
                        for owner in owners:
                            yield [section["title"], row["label"], _OWNER_KINDS[kind],
                                   owner["name"], owner["now"], owner["ref"], owner["delta"]]
                    continue
                if by_owner:
                    for owner in row["owners"] or []:
                        yield [section["title"], row["label"], owner["name"],
                               owner["now"], owner["ref"], owner["delta"]]
                    continue
                v = row["villages"] or {}
                yield [section["title"], row["label"], row["unit"],
                       row["now"], row["ref"], row["delta"],
                       v.get("now"), v.get("ref"), v.get("delta")]

    m = payload["month"]
    return _write([
        ("Month", header, rows()),
        _about([
            ("Month", f"{m['year']}-{m['month']:02d}"),
            ("Compared with", f"{payload['ref_month']['year']}-{payload['ref_month']['month']:02d}"),
            ("Cut at day", payload["day"] or "whole month"),
            ("View", payload["by"]),
            ("Empty cells", "Not recorded (before Mehr 1405)"),
        ]),
    ])


_AREA_MEASURES = (("onair", "On air"), ("dt", "DT done"), ("full", "Fully approved"),
                  ("ict", "ICT approved"), ("cra", "CRA approved"))


def area(payload: dict) -> BuiltWorkbook:
    """One row per breakdown member, then the scope and national on their
    own sheet. On air, DT done and Fully approved are % of villages; ICT and
    CRA are % of DT-done villages."""
    measure_header = [h for _k, label in _AREA_MEASURES for h in (label, f"{label} %")]
    header = ["Name", "Villages", *measure_header, "Compared"]

    def figures(measures: dict) -> list:
        return [v for k, _l in _AREA_MEASURES for v in (measures[k]["count"], measures[k]["rate"])]

    rows = (
        [r["name"], r["villages"], *figures(r["measures"]),
         "Not compared" if r["low_sample"] else "Yes"]
        for r in payload["rows"]
    )
    totals, national = payload["totals"], payload["national"]
    headline = (
        [label, totals[k]["count"], totals[k]["rate"], national[k]]
        for k, label in _AREA_MEASURES
    )
    return _write([
        ("Breakdown", header, rows),
        ("Scope", ["Measure", "Count", "Rate %", "National %"],
         [["Villages", totals["villages"], None, None], *headline]),
        _about([("Scope", payload["scope"]["label"]), ("Breakdown", payload["breakdown"]),
                ("As of", "today")]),
    ])
