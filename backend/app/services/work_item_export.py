"""The assigned-sites spreadsheet.

One export, five columns, deliberately: a contractor asking for "my sites"
wants the list they will drive to — site, type, province, when it landed on
them, and what technology was asked for. Everything else the Work Items table
carries (approval users, stage history, internal ids) is somebody else's
column, and putting it in the file only invites it into the conversation.

Rows come from ``work_item_rows.build_list_row``, so the file and the screen
can never disagree about what a site's assignment date or province is.
"""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="0B8477")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# Header -> the list-row key it reads. The order here is the order in the file.
_COLUMNS: list[tuple[str, str]] = [
    ("Site ID", "site_code"),
    ("Site Type", "site_type"),
    ("Province", "province"),
    ("Assigned Date", "assignment_date"),
    ("Requested Technology", "requested_technology"),
]

# Wide enough that no column opens truncated, narrow enough that the sheet
# fits a laptop screen. openpyxl has no autofit, so this is the substitute.
_WIDTHS = [16, 14, 18, 14, 20]


def _cell(value) -> object:
    """Render one value for a spreadsheet cell.

    Dates go in as dates rather than strings so Excel sorts and filters them
    as dates; a tz-aware datetime is dropped to its date, both because the
    assignment is a day-grained fact and because openpyxl refuses to write
    timezone-aware datetimes at all.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return value if value is not None else ""


def build_sheet(
    title: str, columns: list[tuple[str, str]], widths: list[int], rows: list[dict]
) -> bytes:
    """One .xlsx sheet: a header row from ``columns`` (header, row key), then
    a line per row. Dates stay dates; the header is frozen and filterable."""
    wb = Workbook()
    ws = wb.active
    ws.title = title

    ws.append([header for header, _ in columns])
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = widths[col_idx - 1]

    for row in rows:
        ws.append([_cell(row.get(key)) for _, key in columns])
        for col_idx, (_, key) in enumerate(columns, start=1):
            if isinstance(row.get(key), (date, datetime)):
                ws.cell(row=ws.max_row, column=col_idx).number_format = "yyyy-mm-dd"

    # The header stays put while scrolling, and every column gets a filter —
    # a contractor with four hundred sites reads this file by filtering it.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_assigned_sites_export(rows: list[dict]) -> bytes:
    """Build the .xlsx for a list of work-item rows. Returns the file bytes."""
    return build_sheet("Assigned Sites", _COLUMNS, _WIDTHS, rows)
