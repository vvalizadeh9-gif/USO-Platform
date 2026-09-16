"""The drill-through site list as a spreadsheet.

Same rows, same order and the same columns as the screen, because the two are
read side by side: a reader exports what they are looking at and sends it on,
and a file whose columns differ from the page they came from starts an
argument about which one is right.

Above the table sits a header block naming the filters that were applied, when
the file was generated and how many rows it holds. Without it the file is a
list of sites with no statement of what it is a list *of* -- which is exactly
the file that gets forwarded and misread a week later.
"""
from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import jalali

_HEADER_FILL = PatternFill("solid", fgColor="0B8477")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_META_FONT = Font(color="666666", size=10)

#: Column header -> the row key it reads. The order here is the order on the
#: screen, and both are the order in the file.
COLUMNS: list[tuple[str, str]] = [
    ("Site ID", "site_code"),
    ("Villages", "villages"),
    ("Province", "province"),
    ("Contractor", "contractor"),
    ("Bucket", "bucket"),
    ("Stage", "current_stage"),
    ("Launch date", "launch_date"),
    ("Days since launch", "days_since_launch"),
    ("Assigned", "assignment_date"),
    ("Days held", "days_since_assignment"),
    ("How long", "age_band"),
    ("Problem categories", "problem_categories"),
    ("Fix owners", "fix_owners"),
    ("Oldest open fix (days)", "oldest_open_fix_days"),
    ("Days late", "max_days_late"),
    ("HC round", "hc_round"),
    ("DT execution date", "dt_execution_date"),
    ("DT approved", "dt_approved_at"),
    ("Evidence", "dt_evidence_count"),
]

_WIDTHS = [16, 28, 16, 22, 14, 22, 14, 17, 14, 12, 14, 26, 24, 20, 12, 10, 18, 14, 10]

#: How the header block names each filter, in the order it reads them.
_FILTER_LABELS: list[tuple[str, str]] = [
    ("bucket", "Figure"),
    ("category", "Category"),
    ("age_band", "How long"),
    ("stage", "Stage"),
    ("contractor_id", "Contractor"),
    ("province_id", "Province"),
    ("year", "Shamsi year"),
    ("month", "Shamsi month"),
    ("overdue", "Overdue only"),
    ("owner_role_id", "Fix owner role"),
    ("sort", "Sorted by"),
]


def _cell(value) -> object:
    """Render one value for a spreadsheet cell.

    Lists are joined rather than written as their Python repr, and a missing
    value is an empty cell rather than a zero -- a site with no open fix has
    not waited nought days.
    """
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if value is None:
        return ""
    return value


def build_site_list_export(
    rows: list[dict],
    filters: dict[str, str],
    *,
    generated_on: date | None = None,
) -> bytes:
    """Build the .xlsx for a drill-through list. Returns the file bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Drive Test Sites"

    stamp = jalali.format_shamsi(generated_on or date.today())
    described = describe(filters)

    ws["A1"] = "Drive Test — site list"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = described
    ws["A2"].font = _META_FONT
    ws["A3"] = f"Generated {stamp} · {len(rows)} row{'' if len(rows) == 1 else 's'}"
    ws["A3"].font = _META_FONT
    for row in (1, 2, 3):
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center")

    header_row = 5
    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = _WIDTHS[col_idx - 1]

    for offset, row in enumerate(rows, start=header_row + 1):
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            ws.cell(row=offset, column=col_idx, value=_cell(row.get(key)))

    # The header stays put while scrolling and every column gets a filter: a
    # reader with four hundred rows reads this file by filtering it.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row + len(rows)}"
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def describe(filters: dict[str, str]) -> str:
    """The applied filters as one readable line, for the header block."""
    parts = [
        f"{label}: {filters[key]}"
        for key, label in _FILTER_LABELS
        if filters.get(key) is not None
    ]
    return " · ".join(parts) if parts else "No filters applied"


def filename(filters: dict[str, str], *, generated_on: date | None = None) -> str:
    """A filename that says what the file is: ``dt-problematic-temp-power-1405-06-25.xlsx``.

    Built from the filters rather than from the screen's title so it cannot
    describe a different list from the one inside it. Non-ASCII is dropped
    rather than transliterated -- a Persian contractor name in a filename
    survives none of the round trips these files make through mail clients and
    Windows shares -- which is why the id is kept beside it.
    """
    stamp = (jalali.format_shamsi(generated_on or date.today()) or "").replace("/", "-")
    pieces = ["dt"]
    for key in ("bucket", "category", "age_band", "stage", "contractor_id", "province_id"):
        value = filters.get(key)
        if value:
            pieces.append(str(value))
    slug = _slug("-".join(pieces)) or "dt"
    return f"{slug}-{stamp}.xlsx"


def _slug(text: str) -> str:
    """Lowercase ASCII, hyphen-separated, with everything else stripped."""
    ascii_only = text.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", ascii_only)).strip("-")
