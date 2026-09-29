"""The villages behind a Lifecycle Gaps figure, as a spreadsheet.

``services/gaps.py`` decides *which* villages (:func:`gaps.export_villages`);
this module only writes them. Two sheets:

* **Villages** -- one row per village, with who owns it under every lens, both
  authorities' standing and dates, Mojri's standing and why nobody owns it
  when nobody does. Bold frozen header, a filter on every column.
* **Summary** -- what the file is a list *of*: the figure, the filter, how
  many villages, when and by whom it was exported. A column of village ids
  with no statement of what they are is the file that gets misread a week
  later (see ``acceptance_site_export.py``, which this follows).

The workbook is written in openpyxl's write-only mode, row by row, into a
spooled temporary file: a large export is streamed from the database into the
file and from the file to the client, never held in memory whole.
"""
from __future__ import annotations

import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.core import jalali
from app.services.dt_site_export import cell_value

#: Column header -> the row key it reads, and its width.
COLUMNS: list[tuple[str, str, int]] = [
    ("Village ID", "village_code", 14),
    ("Village name", "village_name", 28),
    ("Province", "province", 18),
    ("CRA region", "region", 16),
    ("Regional manager", "rm", 18),
    ("Coordinator", "coordinator", 18),
    ("Contractor", "contractor", 22),
    ("ICT status", "ict_status", 12),
    ("ICT date", "ict_date", 12),
    ("CRA status", "cra_status", 12),
    ("CRA date", "cra_date", 12),
    ("Mojri status", "mojri", 26),
    ("Attribution", "attribution", 22),
]

#: Keys whose cells are Farsi text, set right-to-left.
_RTL_KEYS = {"village_name"}
_DATE_KEYS = {"ict_date", "cra_date"}

_BOLD = Font(bold=True)
_RTL = Alignment(horizontal="right", readingOrder=2)

#: Kept in memory up to this size, spilled to a temporary file beyond it.
_SPOOL_BYTES = 8 * 1024 * 1024
#: Above this many villages the file is streamed to the client in chunks
#: rather than sent as one body.
STREAM_THRESHOLD_ROWS = 20_000
CHUNK_BYTES = 64 * 1024


@dataclass(frozen=True)
class ExportMeta:
    """What the Summary sheet says about the file."""

    title: str
    filter_text: str
    exported_by: str
    exported_at: datetime


@dataclass
class BuiltWorkbook:
    """A finished workbook in a (spooled) temporary file, rewound, and how
    many villages it holds. The caller closes ``file``."""

    file: tempfile.SpooledTemporaryFile
    rows: int

    def read(self) -> bytes:
        return self.file.read()

    def chunks(self):
        try:
            while chunk := self.file.read(CHUNK_BYTES):
                yield chunk
        finally:
            self.file.close()


def build(rows: Iterable[dict], meta: ExportMeta) -> BuiltWorkbook:
    """Write the Villages and Summary sheets; return the rewound file."""
    wb = Workbook(write_only=True)
    count = _write_villages(wb.create_sheet("Villages"), rows)
    _write_summary(wb.create_sheet("Summary"), meta, count)

    target = tempfile.SpooledTemporaryFile(max_size=_SPOOL_BYTES)
    wb.save(target)
    target.seek(0)
    return BuiltWorkbook(file=target, rows=count)


def _write_villages(ws, rows: Iterable[dict]) -> int:
    # Set before the first row: write-only sheets emit their view and column
    # settings ahead of the data.
    ws.freeze_panes = "A2"
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.append([_header(ws, title) for title, _, _ in COLUMNS])
    count = 0
    for row in rows:
        ws.append([_cell(ws, key, row.get(key)) for _, key, _ in COLUMNS])
        count += 1

    # The filter covers the header and every row: a reader with a thousand
    # villages reads this file by filtering it.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{count + 1}"
    return count


def _header(ws, title: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=title)
    cell.font = _BOLD
    return cell


def _cell(ws, key: str, value):
    if key in _DATE_KEYS:
        return jalali.format_shamsi(value) if value else None
    if key in _RTL_KEYS and value:
        cell = WriteOnlyCell(ws, value=cell_value(value))
        cell.alignment = _RTL
        return cell
    return cell_value(value)


def _write_summary(ws, meta: ExportMeta, count: int) -> None:
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60
    when = meta.exported_at.astimezone(timezone.utc)
    lines = [
        ("Figure", meta.title),
        ("Filter", meta.filter_text),
        ("Villages", count),
        ("Exported (Jalali)", jalali.format_shamsi(jalali.tehran_today())),
        ("Exported (Gregorian)", when.strftime("%Y-%m-%d %H:%M UTC")),
        ("Exported by", meta.exported_by),
    ]
    for label, value in lines:
        ws.append([_header(ws, label), value])


def filename(gap: str, file_key: str, *, generated_on: date | None = None) -> str:
    """``uep-pending-cra-coordinator-v-hashemi-1405-07-07.xlsx``: ASCII only,
    built from the request rather than from screen wording, so it cannot
    describe a different list from the one inside."""
    stamp = (jalali.format_shamsi(generated_on or jalali.tehran_today()) or "").replace("/", "-")
    return f"uep-{gap.replace('_', '-')}-{file_key}-{stamp}.xlsx"
