"""The Drive Test delivery workbook.

One file that holds the whole DT picture for the caller's scope: what the
dashboard says, the sites behind it, the fixes that are open, the contractor
ledger and the provinces. It exists because the audiences for those five
things -- an audit, a management meeting, the regulator -- read them together
and offline, and until now that meant five screenshots and a promise that they
were taken at the same moment.

**Every figure comes from an existing service call.** ``DriveTestAnalytics``
for the KPIs and breakdowns, ``dt_site_list`` for the site rows,
``DriveTestAnalytics.scorecard`` for the ledger, ``hc_queues.remediations``
for the open fixes. Nothing is counted a second time here, because a workbook
that computes its own version of a dashboard figure is a workbook that will
one day disagree with the dashboard, in a meeting, with no way to tell which
one is right.

What is deliberately absent: anything from the monthly snapshots -- the trend
series and the month's flow ledger. Those are a different vintage of data (see
``api/drive_test.drive_test_trend``), and a file whose figures were taken at
one moment should not carry figures taken at another. The Summary sheet says
so in as many words rather than leaving the omission to be noticed.

**Privacy.** Scope decides the rows before this module sees them, and nothing
here widens it. A contractor account's file carries its own sites, its own row
in the ledger and no other company's name anywhere -- the test for that reads
every cell of every sheet rather than trusting this paragraph. A category
owner sees the sites routed to them and their own fixes, and gets no ledger
sheet at all: the ledger is a comparison between companies, which is not their
business and not their scope.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core import jalali
from app.services import dt_site_export, dt_site_list, hc_queues, pip_export
from app.services import monthly_plan as plans
from app.services.drive_test_analytics import DriveTestAnalytics, own_contractor_id

#: How many Shamsi months the contractor ledger covers.
LEDGER_MONTHS = 6

_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="0B8477")
_TITLE = Font(bold=True, size=14)
_SECTION = Font(bold=True, size=11)
_NOTE = Font(italic=True, size=9, color="55637A")
_LABEL = Font(color="55637A")

#: Right-to-left, for the columns that carry Persian text.
#:
#: ``readingOrder=2`` as well as right alignment: without it Excel lays a
#: Persian string out left-to-right and a name that ends in a bracket or a
#: digit comes apart on screen. The dashboard solves the same problem with the
#: ``dt-farsi`` class.
_RTL = Alignment(horizontal="right", readingOrder=2)

#: Which of the site columns carry Persian text.
_FARSI_KEYS = {"villages", "province", "contractor"}


class WorkbookTooLarge(Exception):
    """A sheet would exceed the export cap.

    Raised rather than truncating. A file that silently stops at twenty
    thousand rows is read as the whole answer, which is worse than no file --
    and this one is built for audits.
    """


def _shamsi_now() -> str:
    """The generation stamp: Shamsi date and clock time, UTC."""
    now = datetime.now(timezone.utc)
    return f"{jalali.format_shamsi(now.date())} {now:%H:%M} UTC"


def filename(scope: str | None, *, generated_on: date | None = None) -> str:
    """``dt-delivery-kerman-1405-06-25.xlsx``, ASCII-safe.

    Non-ASCII is dropped rather than transliterated, so a Persian province
    name leaves ``dt-delivery-1405-06-25.xlsx`` -- which is honest about the
    file being scoped without inventing a spelling of the scope that nothing
    else in the platform uses. The scope is named inside the file, on the
    Summary sheet, where it cannot be lost by a mail client.
    """
    stamp = (jalali.format_shamsi(generated_on or date.today()) or "").replace("/", "-")
    slug = _slug(scope or "")
    return f"dt-delivery-{slug + '-' if slug else ''}{stamp}.xlsx"


def _slug(text: str) -> str:
    ascii_only = text.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", ascii_only)).strip("-")


def build(
    db: Session,
    user,
    *,
    province_id: int | None = None,
    province_name: str | None = None,
    scope_label: str | None = None,
    max_rows: int,
) -> bytes:
    """The delivery workbook for this caller's scope, as .xlsx bytes.

    ``province_id`` is applied where the dashboard applies it: inside
    ``DriveTestAnalytics``, after ``apply_work_item_scope``, so it can only
    narrow what the caller was already entitled to see.

    ``province_name`` filters the open fixes to the same province, and an
    empty string means "match no province" -- which is what an id outside the
    caller's scope resolves to. That id yields an empty but perfectly valid
    workbook rather than an error, and ``scope_label`` names it by id rather
    than by name, so nothing in the file says whether the province exists.
    The dashboard answers a 404 there instead; this is a file, and a file that
    opens and is empty is a clearer answer than a download that failed.

    Raises :class:`WorkbookTooLarge` if any sheet would exceed *max_rows*.
    """
    analytics = DriveTestAnalytics(db, user, province_id=province_id)
    kpis = analytics.compute_kpis()
    breakdowns = analytics.breakdowns()
    year, month = jalali.current_shamsi_period()
    plan = analytics.plan_and_delivery(year, month)

    filters = dt_site_list.parse_filters(db, user, bucket="onair", province_id=province_id)
    sites = dt_site_list.build_rows(db, user, filters)

    fixes = _open_fixes(db, user, province_name)

    is_contractor = own_contractor_id(user) is not None
    is_category_owner = bool(getattr(getattr(user, "role", None), "is_category_owner", False))
    ledger = (
        None
        if is_category_owner
        else analytics.scorecard(plans.trailing_periods(LEDGER_MONTHS))
    )

    for name, rows in (("Sites", sites), ("Open fixes", fixes)):
        if len(rows) > max_rows:
            raise WorkbookTooLarge(
                f"The {name} sheet would hold {len(rows)} rows, more than the "
                f"{max_rows} this export carries. Narrow to one province and "
                "try again."
            )

    wb = Workbook()
    _summary_sheet(
        wb.active,
        user=user,
        province_name=scope_label,
        kpis=kpis,
        breakdowns=breakdowns,
        plan=plan,
        is_contractor=is_contractor,
    )
    _sites_sheet(wb.create_sheet("Sites"), sites)
    _fixes_sheet(wb.create_sheet("Open fixes"), fixes)
    if ledger is not None:
        pip_export.write_contractor_sheet(wb.create_sheet("Contractor ledger"), ledger["months"])
    _provinces_sheet(wb.create_sheet("Provinces"), breakdowns["provinces"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _open_fixes(db: Session, user, province_name: str | None) -> list[dict]:
    """The open fixes in scope, from the queue service that already answers this.

    Narrowed to one province by name when the file is: province names are
    unique in the schema, and the rows this service returns carry the name
    rather than the id. Filtering here rather than asking for a second,
    differently-scoped query keeps the fixes in the file to the ones whose
    sites are in the file.
    """
    rows = hc_queues.remediations(db, user)
    if province_name is None:
        return rows
    return [r for r in rows if r["province"] == province_name]





# ------------------------------------------------------------------ sheets
def _summary_sheet(ws, *, user, province_name, kpis, breakdowns, plan, is_contractor) -> None:
    """What the dashboard says, at the moment the file was generated."""
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12

    row = 1
    ws.cell(row=row, column=1, value="Drive Test delivery").font = _TITLE
    row += 2

    # Who and when, and for what. The role rather than the person: the file
    # travels, and which account pressed the button is in the platform's own
    # log, not on a sheet that will be forwarded.
    for label, value in (
        ("Generated", _shamsi_now()),
        ("Generated by", getattr(getattr(user, "role", None), "name", "—")),
        ("Scope", province_name or "All provinces in your scope"),
    ):
        ws.cell(row=row, column=1, value=label).font = _LABEL
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    row = _block(
        ws, row, "Programme",
        [
            ("On-air", kpis["total_onair"]),
            ("Drive tests done", kpis["total_dt_done"]),
            ("Ongoing", kpis["total_ongoing"]),
            ("Problematic", kpis["total_problematic"]),
            ("Not started", kpis["total_not_started"]),
            ("Remaining", kpis["total_remaining"]),
        ],
    )

    row = _block(
        ws, row, "Problematic by category",
        [(p["name"], p["value"]) for p in breakdowns["problematic"]["by_category"]],
    )

    ongoing = breakdowns["ongoing"]
    row = _block(
        ws, row, "Ongoing by how long it has been held",
        [(p["name"], p["value"]) for p in ongoing["by_age"]]
        # Reported beside the bands rather than folded into the newest one: a
        # site nobody has been assigned has no clock running on it, and filing
        # it under "up to 1 week" would make the backlog look fresher than it
        # is. The bands run on the assignment date; see ``AGE_BANDS``.
        + [("No assignment date", ongoing["without_assignment_date"])],
    )

    row = _block(
        ws, row, "Ongoing by stage",
        [(p["name"], p["value"]) for p in ongoing["by_stage"]],
    )

    row = _block(
        ws, row, f"Plan and delivery — {plan['month_label']}",
        [
            # "Assignment", not "Assigned": the stage block above already has a
            # row called "Assigned" -- the workflow stage, whose spelling is
            # fixed by ``workflow.py`` and travels in URLs -- and two rows
            # under one label in one column is a lookup down column A that
            # silently answers with the other figure. "Assignment" is also
            # what the Plan and delivery card calls it on screen, which is the
            # property this whole sheet exists to hold.
            ("Assignment", plan["assigned"]),
            ("PIP", plan["pip"]),
            ("Delivered", plan["actual"]),
            ("Achievement %", plan["achievement_percent"]),
        ],
    )

    ws.cell(
        row=row,
        column=1,
        value=(
            "Trend and month-on-month movement are not in this file. Those come "
            "from the monthly snapshots, which are written on a different "
            "schedule from these figures; carrying both would put two vintages "
            "of data in one workbook."
        ),
    ).font = _NOTE
    row += 1
    if is_contractor:
        ws.cell(
            row=row,
            column=1,
            value="Every figure in this file is your own company's work only.",
        ).font = _NOTE


def _block(ws, row: int, title: str, pairs: list[tuple[str, object]]) -> int:
    """A titled label/value block, returning the next free row."""
    ws.cell(row=row, column=1, value=title).font = _SECTION
    row += 1
    for label, value in pairs:
        ws.cell(row=row, column=1, value=label).alignment = _RTL if _is_farsi(label) else None
        # Numbers go in as numbers. A figure stored as text cannot be summed,
        # sorted or charted by the person who opens the file, which is most of
        # what they opened it for.
        ws.cell(row=row, column=2, value=value if value is not None else "—")
        row += 1
    return row + 1


def _is_farsi(text: object) -> bool:
    return isinstance(text, str) and any("؀" <= ch <= "ۿ" for ch in text)


def _sites_sheet(ws, rows: list[dict]) -> None:
    """One row per on-air site, in Prompt 1's columns and order.

    The column list is imported from the site-list export rather than restated,
    so the sheet a reader gets here and the file they get from the drill-through
    screen cannot drift into two different shapes.
    """
    columns = dt_site_export.COLUMNS
    _header(ws, [h for h, _ in columns], dt_site_export.WIDTHS)

    for r, row in enumerate(rows, start=2):
        for c, (_, key) in enumerate(columns, start=1):
            # Through the drill-through export's own renderer, so the two
            # files agree on what an empty cell means -- and so the control
            # characters CPM-imported text carries are stripped before they
            # fail the whole workbook. See ``dt_site_export.cell_value``.
            cell = ws.cell(row=r, column=c, value=dt_site_export.cell_value(row.get(key)))
            if key in _FARSI_KEYS:
                cell.alignment = _RTL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


#: The open-fixes sheet, as ``(header, width, row key)``.
_FIX_COLUMNS = [
    ("Site ID", 16, "site_code"),
    ("Province", 16, "province"),
    ("Category", 24, "category"),
    ("Owner role", 22, "owner_role"),
    ("HC round", 10, "round_no"),
    ("Days open", 11, "days_open"),
    ("Due", 14, "due_shamsi"),
    ("Days late", 11, "days_late"),
    ("Re-route pending", 17, "reroute"),
]


def _fixes_sheet(ws, rows: list[dict]) -> None:
    """Every open fix in scope, worst-late first.

    The order is the queue service's own -- it sorts by days late and then by
    days open -- rather than a second sort applied here, so the sheet and the
    Fix Queue screen put the same row at the top.
    """
    _header(ws, [h for h, _, _ in _FIX_COLUMNS], [w for _, w, _ in _FIX_COLUMNS])

    for r, row in enumerate(rows, start=2):
        due = row.get("due_at")
        values = {
            **row,
            "due_shamsi": jalali.format_shamsi(due.date()) if due else "",
            "reroute": "Yes" if row.get("reroute_pending") else "No",
        }
        for c, (_, _, key) in enumerate(_FIX_COLUMNS, start=1):
            cell = ws.cell(
                row=r, column=c, value=dt_site_export.cell_value(values.get(key))
            )
            if key == "province":
                cell.alignment = _RTL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_FIX_COLUMNS))}{len(rows) + 1}"


_PROVINCE_COLUMNS = [
    ("Province", 20, "name"),
    ("On-air", 10, "onair"),
    ("Done", 10, "done"),
    ("Remaining", 12, "remaining"),
    ("Ongoing", 10, "ongoing"),
    ("Problematic", 13, "problematic"),
    ("Done %", 10, "done_percent"),
]


def _provinces_sheet(ws, rows: list[dict]) -> None:
    """The dashboard's province table, in the order it already sorts itself."""
    _header(ws, [h for h, _, _ in _PROVINCE_COLUMNS], [w for _, w, _ in _PROVINCE_COLUMNS])
    for r, row in enumerate(rows, start=2):
        for c, (_, _, key) in enumerate(_PROVINCE_COLUMNS, start=1):
            cell = ws.cell(
                row=r, column=c, value=dt_site_export.cell_value(row.get(key))
            )
            if key == "name":
                cell.alignment = _RTL
    ws.freeze_panes = "A2"


def _header(ws, titles: list[str], widths: list[int]) -> None:
    for i, title in enumerate(titles, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
