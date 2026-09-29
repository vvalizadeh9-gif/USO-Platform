"""The Acceptance Dashboard's Export: its progress payload as a spreadsheet.

Built from the dict ``acceptance_progress.progress`` returns -- the same call
the page makes -- so the file cannot hold a figure the chart does not. One
sheet per stream (Village, ICT, CRA), one row per month, and the plan columns
only where the caller may see them: a contractor's file has no Internal PIP
column at all, rather than an empty one.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import jalali
from app.services.acceptance_progress import STREAM_KEYS

_HEADER_FILL = PatternFill("solid", fgColor="2F5FD0")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_META_FONT = Font(color="5F6B7E", size=10)

SHEET_TITLES = {"village": "Village", "ict": "ICT", "cra": "CRA"}
STREAM_DESCRIPTIONS = {
    "village": "Villages fully accepted (ICT and CRA approved)",
    "ict": "Villages approved by ICT",
    "cra": "Villages approved by CRA",
}

_APPROVED_COLUMNS = [
    ("Approved", "approved"),
    ("Approved to date", "approved_cumulative"),
]
_INTERNAL_COLUMNS = [
    ("Internal PIP", "internal_plan"),
    ("Internal PIP to date", "internal_plan_cumulative"),
]
_CONTRACTOR_COLUMNS = [
    ("Contractor PIP", "contractor_plan"),
    ("Contractor PIP to date", "contractor_plan_cumulative"),
]


def _columns(data: dict) -> list[tuple[str, str]]:
    columns = list(_APPROVED_COLUMNS)
    if data["plans_available"] and data["internal_visible"]:
        columns += _INTERNAL_COLUMNS
    if data["plans_available"]:
        columns += _CONTRACTOR_COLUMNS
    return columns


def _write_sheet(ws, key: str, data: dict, stamp: str) -> None:
    ws["A1"] = f"Acceptance — {STREAM_DESCRIPTIONS[key]}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Generated {stamp} · {len(data['months'])} months · month in progress marked *"
    ws["A2"].font = _META_FONT

    columns = _columns(data)
    header_row = 4
    for col, header in enumerate(["Month", *[h for h, _ in columns]], start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col)].width = 14 if col == 1 else 20

    for offset, month in enumerate(data["months"], start=header_row + 1):
        label = f"{month['label']} {month['shamsi_year']}" + (" *" if month["is_current"] else "")
        ws.cell(row=offset, column=1, value=label)
        for col, (_, field) in enumerate(columns, start=2):
            ws.cell(row=offset, column=col, value=month[key][field])
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)


def build_progress_export(data: dict, *, generated_on: date | None = None) -> bytes:
    """Build the .xlsx for a progress payload. Returns the file bytes."""
    stamp = jalali.format_shamsi(generated_on or jalali.tehran_today()) or ""
    wb = Workbook()
    for index, key in enumerate(STREAM_KEYS):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = SHEET_TITLES[key]
        _write_sheet(ws, key, data, stamp)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(*, generated_on: date | None = None) -> str:
    """``acceptance-progress-1405-07-07.xlsx``."""
    stamp = (jalali.format_shamsi(generated_on or jalali.tehran_today()) or "").replace("/", "-")
    return f"acceptance-progress-{stamp}.xlsx"
