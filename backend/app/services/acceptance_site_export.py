"""The sites behind an Acceptance dashboard figure, as a spreadsheet.

The same rows, in the same order, with the same columns as the drill-through
panel on screen — because the two are read side by side. A reader exports what
they are looking at and forwards it, and a file whose columns differ from the
page it came from starts an argument about which one is right.

Above the table sits a header block naming the figure the list is of, when the
file was generated, how many sites it holds and how many villages those sites
account for. Without it the file is a column of site ids with no statement of
what they are a list *of*, which is exactly the file that gets misread a week
later.

The cell rendering is ``dt_site_export.cell_value`` rather than a second copy
of it: these rows come out of the same CPM workbooks, carry the same stray
control characters, and one file quietly disagreeing with the other about what
an empty cell means is the failure that module already exists to prevent.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import jalali
from app.services.dt_site_export import cell_value

_HEADER_FILL = PatternFill("solid", fgColor="0B8477")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_META_FONT = Font(color="666666", size=10)

#: Column header -> the row key it reads. The order here is the order in the
#: panel, and both are the order in the file. Site ID leads because it is what
#: the file is opened to find.
COLUMNS: list[tuple[str, str]] = [
    ("Site ID", "site_code"),
    ("Province", "province"),
    ("Villages", "villages"),
    ("Village names", "village_names"),
    ("DT done", "dt_done"),
    ("Approved", "approved"),
    ("Rejected", "rejected"),
    ("Waiting", "pending"),
]

WIDTHS = [18, 18, 10, 44, 10, 11, 11, 10]


def build_site_list_export(
    rows: list[dict],
    *,
    label: str,
    total: int,
    generated_on: date | None = None,
) -> bytes:
    """Build the .xlsx for one figure's site list. Returns the file bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Acceptance Sites"

    stamp = jalali.format_shamsi(generated_on or date.today())
    sites = len(rows)

    ws["A1"] = "Acceptance — sites behind a figure"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Figure: {label}"
    ws["A2"].font = _META_FONT
    ws["A3"] = (
        f"Generated {stamp} · {sites} site{'' if sites == 1 else 's'} · "
        f"{total} village row{'' if total == 1 else 's'}"
    )
    ws["A3"].font = _META_FONT
    for row in (1, 2, 3):
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center")

    header_row = 5
    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = WIDTHS[col_idx - 1]

    for offset, row in enumerate(rows, start=header_row + 1):
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            ws.cell(row=offset, column=col_idx, value=cell_value(row.get(key)))

    # The header stays put while scrolling and every column gets a filter: a
    # reader with four hundred sites reads this file by filtering it.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row + sites}"
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(metric: str, *, generated_on: date | None = None) -> str:
    """``acceptance-rejected-1405-06-25.xlsx`` — a name that says what is inside.

    Built from the metric rather than from the screen's wording so it cannot
    describe a different list from the one in the file.
    """
    stamp = (jalali.format_shamsi(generated_on or date.today()) or "").replace("/", "-")
    slug = metric.replace("_", "-")
    return f"acceptance-{slug}-{stamp}.xlsx"
