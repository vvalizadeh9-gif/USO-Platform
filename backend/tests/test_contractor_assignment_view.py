"""The contractor's view of an assignment: KPI, detail fields, and export.

Three things a subcontractor asked for and the screen did not answer:

* how many sites am I holding right now (a counter, not a wall of cards),
* how long has this one been sitting with me (aging measured from the
  assignment, which runs while the work is still open),
* give me the list as a spreadsheet — and only mine.

Run with:  cd backend && pytest tests/test_contractor_assignment_view.py -q
"""
import io
import os
import sys
from datetime import date, datetime, timedelta, timezone

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_contractor_view_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from tests.conftest import create_schema, login_form  # noqa: E402


@pytest.fixture(scope="module")
def client():
    if os.path.exists("/tmp/uep_contractor_view_pytest.db"):
        os.remove("/tmp/uep_contractor_view_pytest.db")
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        yield c


def _login(client, username="admin", password="Admin@12345") -> dict:
    r = client.post("/api/v1/auth/login", data=login_form(client, username, password))
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _seed_assigned_site(*, contractor_name: str, site_code: str, days_ago: int) -> dict:
    """One assigned site held by a (possibly new) contractor."""
    from app.models.reference import Contractor, Province
    from app.models.workitem import Assignment, Site, WorkItem

    db = SessionLocal()
    province = (
        db.query(Province).filter(Province.name == "Kerman").one_or_none()
        or Province(name="Kerman")
    )
    contractor = (
        db.query(Contractor).filter(Contractor.name == contractor_name).one_or_none()
        or Contractor(name=contractor_name, type="drive_test", active=True)
    )
    db.add_all([province, contractor])
    db.flush()

    site = Site(site_code=site_code, province_id=province.id)
    db.add(site)
    db.flush()

    wi = WorkItem(
        site_id=site.id,
        site_type="Rural",
        last_stage="راه_اندازی_دائم",
        requested_technology="2G/3G/4G",
        current_stage="Assigned",
    )
    db.add(wi)
    db.flush()
    wi.assignments.append(
        Assignment(
            assignment_type="official",
            contractor_id=contractor.id,
            assigned_at=datetime.now(timezone.utc) - timedelta(days=days_ago),
            is_active=True,
        )
    )
    db.commit()
    out = {"work_item_id": wi.id, "contractor_id": contractor.id}
    db.close()
    return out


def _make_contractor_user(client, admin_h, *, username: str, contractor_id: int) -> dict:
    roles = client.get("/api/v1/reference/roles", headers=admin_h).json()
    role_id = next(r["id"] for r in roles if r["name"] == "Contractor")
    r = client.post(
        "/api/v1/admin/users",
        headers=admin_h,
        json={
            "username": username,
            "password": "Test-Fixture-Passphrase",
            "first_name": "Field", "family_name": username,
            "role_id": role_id,
            "contractor_id": contractor_id,
            "sees_all_provinces": False,
            "province_ids": [],
        },
    )
    assert r.status_code == 201, r.text
    return _login(client, username, "Test-Fixture-Passphrase")


@pytest.fixture(scope="module")
def world(client):
    """Two contractors, one assigned site each, and a login for the first."""
    admin_h = _login(client)
    mine = _seed_assigned_site(
        contractor_name="Alpha DT", site_code="CV-MINE-1", days_ago=9
    )
    theirs = _seed_assigned_site(
        contractor_name="Beta DT", site_code="CV-THEIRS-1", days_ago=2
    )
    contractor_h = _make_contractor_user(
        client, admin_h, username="cv_alpha_user", contractor_id=mine["contractor_id"]
    )
    return {
        "admin": admin_h,
        "contractor": contractor_h,
        "mine": mine,
        "theirs": theirs,
    }


# ---------------------------------------------------------------------------
# Aging measured from the assignment
# ---------------------------------------------------------------------------
def test_aging_since_assignment_runs_while_the_work_is_open():
    """An open site ages every day; the closed-interval aging stays None."""
    from app.services.work_item_rows import _aging_days, aging_since_assignment

    class _A:
        assigned_at = datetime.now(timezone.utc) - timedelta(days=12)

    assert _aging_days(_A(), None) is None  # nothing to measure to, yet
    assert aging_since_assignment(_A(), None) == 12


def test_aging_since_assignment_stops_at_the_drive_test():
    from app.services.work_item_rows import aging_since_assignment

    class _A:
        assigned_at = datetime.now(timezone.utc) - timedelta(days=12)

    class _D:
        execution_date = date.today() - timedelta(days=4)

    assert aging_since_assignment(_A(), _D()) == 8


def test_aging_since_assignment_is_never_negative_and_tolerates_nothing():
    from app.services.work_item_rows import aging_since_assignment

    class _A:
        assigned_at = datetime.now(timezone.utc)

    class _D:
        execution_date = date.today() - timedelta(days=10)

    assert aging_since_assignment(_A(), _D()) == 0
    assert aging_since_assignment(None, _D()) is None


# ---------------------------------------------------------------------------
# The detail a subcontractor is shown
# ---------------------------------------------------------------------------
def test_detail_carries_the_five_fields_the_contractor_needs(client, world):
    detail = client.get(
        f"/api/v1/work-items/{world['mine']['work_item_id']}",
        headers=world["contractor"],
    )
    assert detail.status_code == 200, detail.text
    body = detail.json()

    assert body["site_code"] == "CV-MINE-1"
    assert body["site_type"] == "Rural"
    assert body["province"] == "Kerman"
    assert body["assignment_date"] is not None
    assert body["assigned_aging_days"] == 9


def test_detail_no_longer_carries_villages(client, world):
    """Acceptance is its own process — it does not belong on an assignment."""
    body = client.get(
        f"/api/v1/work-items/{world['mine']['work_item_id']}",
        headers=world["admin"],
    ).json()
    assert "villages" not in body


# ---------------------------------------------------------------------------
# The KPI card
# ---------------------------------------------------------------------------
def test_contractor_gets_an_assigned_sites_counter(client, world):
    summary = client.get(
        "/api/v1/action-center/summary", headers=world["contractor"]
    ).json()
    counter = next(
        (c for c in summary["counters"] if c["key"] == "assigned_sites"), None
    )
    assert counter is not None, "a contractor must be told how many sites they hold"
    assert counter["count"] == 1  # theirs, never Beta's
    assert counter["url"] == "/work-items?stage=Assigned"


def test_the_counter_counts_only_this_contractors_sites(client, world):
    """A second site for the same contractor moves the number by exactly one."""
    _seed_assigned_site(
        contractor_name="Alpha DT", site_code="CV-MINE-2", days_ago=1
    )
    summary = client.get(
        "/api/v1/action-center/summary", headers=world["contractor"]
    ).json()
    counter = next(c for c in summary["counters"] if c["key"] == "assigned_sites")
    assert counter["count"] == 2


# ---------------------------------------------------------------------------
# The export
# ---------------------------------------------------------------------------
def _sheet(response):
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content)).active


def test_export_has_exactly_the_five_requested_columns(client, world):
    ws = _sheet(
        client.get(
            "/api/v1/work-items/export",
            headers=world["contractor"],
            params={"stage": "Assigned"},
        )
    )
    header = [c.value for c in ws[1]]
    assert header == [
        "Site ID",
        "Site Type",
        "Province",
        "Assigned Date",
        "Requested Technology",
    ]

    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"CV-MINE-1", "CV-MINE-2"}
    mine = rows["CV-MINE-1"]
    assert mine[1] == "Rural"
    assert mine[2] == "Kerman"
    assert mine[3] is not None
    assert mine[4] == "2G/3G/4G"


def test_export_never_leaks_another_contractors_sites(client, world):
    """The export is scoped exactly like the list it comes from."""
    ws = _sheet(
        client.get("/api/v1/work-items/export", headers=world["contractor"])
    )
    codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert "CV-THEIRS-1" not in codes

    admin_ws = _sheet(
        client.get("/api/v1/work-items/export", headers=world["admin"])
    )
    admin_codes = {row[0] for row in admin_ws.iter_rows(min_row=2, values_only=True)}
    assert {"CV-MINE-1", "CV-THEIRS-1"} <= admin_codes


def test_export_route_is_not_read_as_a_work_item_id(client, world):
    """`/work-items/export` must not fall through to `/work-items/{id}`."""
    r = client.get("/api/v1/work-items/export", headers=world["admin"])
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
