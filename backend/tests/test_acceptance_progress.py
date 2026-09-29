"""Tests for ``GET /acceptance/progress``: approvals against both plans.

Covers the rules in ``services/acceptance_progress.py``:

1. Monthly counts land in the Shamsi month the stream cleared in, and an
   undated approval is an opening balance -- in every month's running total,
   in no month's count.
2. The last month's running total equals the overview's approved figure for
   the same scope, for all three streams.
3. Plans are None when there is none; cumulative plans are anchored on
   actuals; Contractor PIP is the sum across contractors, or one's own.
4. A contractor never receives the Internal PIP and always gets its own PIP.
5. A province scope has no plans at all.
6. ``today``, ``is_current`` and ``days_in_month`` across 31/30/29-day months.
7. The ICT and CRA plan streams round-trip through the Monthly Plan API and
   reach the chart.

The service tests pin ``today`` to 7 مهر 1405 and seed their own months, so
they do not depend on when the suite runs.

Run with:  cd backend && pytest tests/test_acceptance_progress.py -q
"""
import os
import sys

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_acc_progress_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402

from app.core import jalali  # noqa: E402
from app.core.database import SessionLocal  # noqa: E402
from tests.conftest import create_schema, login_as_role, login_form  # noqa: E402

ACC = "/api/v1/acceptance"
PIP = "/api/v1/pip"

TODAY = jalali.from_shamsi_date(1405, 7, 7)
STREAMS = ("village", "ict", "cra")


@pytest.fixture(scope="module")
def client():
    if os.path.exists("/tmp/uep_acc_progress_pytest.db"):
        os.remove("/tmp/uep_acc_progress_pytest.db")
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        yield c


def _login(client, username="admin", password="Admin@12345"):
    r = client.post("/api/v1/auth/login", data=login_form(client, username, password))
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


# ---------------------------------------------------------------------------
# Seed
# ---------------------------------------------------------------------------
def _d(year, month, day=10):
    return jalali.from_shamsi_date(year, month, day)


def _village(db, site, contractor_id, code, *, ict=None, cra=None):
    """One DT-done هدف village on its own work item. ``ict`` / ``cra`` are
    (status, date) or None for nothing filed."""
    from app.models.acceptance import Acceptance
    from app.models.workitem import Village, WorkItem

    wi = WorkItem(
        site_id=site.id, site_type=code, dt_status="Done",
        requested_technology="2G", current_stage="New",
        dt_sc_contractor_id=contractor_id,
    )
    db.add(wi)
    db.flush()
    acc = Acceptance(technology="2G")
    if ict:
        acc.ict_status, acc.ict_date = ict
    if cra:
        acc.cra_status, acc.cra_date = cra
    village = Village(work_item_id=wi.id, village_code=code, target_classification="هدف")
    village.acceptances = [acc]
    db.add(village)


def _pip(db, contractor_id, stream, period, count, version=1, status="Approved"):
    from app.models.monthly_plan import ContractorMonthlyPlan

    db.add(
        ContractorMonthlyPlan(
            contractor_id=contractor_id, stream=stream,
            shamsi_year=period[0], shamsi_month=period[1],
            version=version, is_current=True, committed_count=count, status=status,
        )
    )


@pytest.fixture(scope="module")
def world(client):
    """Two provinces, two contractors, six villages, and some plans.

    ======  ==========  ======================  ======================
    code    contractor  ICT                     CRA
    ======  ==========  ======================  ======================
    V1      C1, P1      Approved 1405-05        Approved 1405-06
    V2      C1, P1      Approved 1405-07        --
    V3      C2, P2      Approved, no date       Approved, no date
    V4      C2, P2      Rejected                Approved 1405-03
    V5      C1, P1      --                      --
    V6      C2, P1      Approved 1403-01        Approved 1403-01
    ======  ==========  ======================  ======================
    """
    from app.models.reference import Contractor, Province, User
    from app.models.workitem import Site
    from app.services import acceptance_plan

    admin_h = _login(client)
    token = admin_h["Authorization"].split()[1]
    pm_h = login_as_role(client, token, "PM")
    coord_h = login_as_role(client, token, "Coordinator")

    db = SessionLocal()
    p1, p2 = Province(name="ProgP1"), Province(name="ProgP2")
    c1 = Contractor(name="Prog C1", type="drive_test", active=True)
    c2 = Contractor(name="Prog C2", type="drive_test", active=True)
    db.add_all([p1, p2, c1, c2])
    db.flush()
    s1, s2 = Site(site_code="PG-1", province_id=p1.id), Site(site_code="PG-2", province_id=p2.id)
    db.add_all([s1, s2])
    db.flush()

    ok, no = "Approved", "Rejected"
    _village(db, s1, c1.id, "V1", ict=(ok, _d(1405, 5)), cra=(ok, _d(1405, 6)))
    _village(db, s1, c1.id, "V2", ict=(ok, _d(1405, 7, 2)))
    _village(db, s2, c2.id, "V3", ict=(ok, None), cra=(ok, None))
    _village(db, s2, c2.id, "V4", ict=(no, _d(1405, 2)), cra=(ok, _d(1405, 3, 1)))
    _village(db, s1, c1.id, "V5")
    _village(db, s1, c2.id, "V6", ict=(ok, _d(1403, 1, 5)), cra=(ok, _d(1403, 1, 5)))

    # Internal PIP (ACCEPTANCE) for شهریور and مهر.
    pm = db.query(User).filter_by(username="test_pm").one()
    acceptance_plan.set_target(db, year=1405, month=6, target_count=5, user=pm)
    acceptance_plan.set_target(db, year=1405, month=7, target_count=4, user=pm)
    # Contractor PIPs for مهر. C1's second approved version is the one in
    # force; a submitted (not approved) PIP never counts.
    _pip(db, c1.id, "ACCEPTANCE", (1405, 7), 2, version=1)
    _pip(db, c1.id, "ACCEPTANCE", (1405, 7), 3, version=2)
    _pip(db, c2.id, "ACCEPTANCE", (1405, 7), 1)
    _pip(db, c2.id, "ACCEPTANCE", (1405, 6), 9, status="Submitted")
    db.commit()

    contractor_h = _contractor_user(client, admin_h, c1.id)
    ids = {"p1": p1.id, "p2": p2.id, "c1": c1.id, "c2": c2.id}
    db.close()
    return {
        **ids,
        "admin": admin_h, "pm": pm_h, "coordinator": coord_h, "contractor": contractor_h,
    }


def _contractor_user(client, admin_h, contractor_id):
    roles = client.get("/api/v1/reference/roles", headers=admin_h).json()
    role_id = next(r["id"] for r in roles if r["name"] == "Contractor")
    password = "Test-Fixture-Passphrase"
    r = client.post(
        "/api/v1/admin/users",
        headers=admin_h,
        json={
            "username": "progress_contractor", "password": password,
            "first_name": "Test", "family_name": "Contractor",
            "role_id": role_id, "sees_all_provinces": True, "province_ids": [],
            "contractor_id": contractor_id,
        },
    )
    assert r.status_code == 201, r.text
    return _login(client, "progress_contractor", password)


def _progress(username, **kwargs):
    """The service's payload for a user, with today pinned, keyed by month."""
    from app.models.reference import User
    from app.services import acceptance_progress

    db = SessionLocal()
    try:
        user = db.query(User).filter_by(username=username).one()
        data = acceptance_progress.progress(db, user, today=kwargs.pop("today", TODAY), **kwargs)
    finally:
        db.close()
    data["by"] = {(m["shamsi_year"], m["shamsi_month"]): m for m in data["months"]}
    return data


# ---------------------------------------------------------------------------
# 1. Monthly counts and the opening balance
# ---------------------------------------------------------------------------
def test_monthly_counts_land_in_the_month_the_stream_cleared(world):
    by = _progress("test_pm")["by"]
    assert by[(1405, 5)]["ict"]["approved"] == 1       # V1
    assert by[(1405, 7)]["ict"]["approved"] == 1       # V2
    assert by[(1405, 6)]["cra"]["approved"] == 1       # V1
    assert by[(1405, 3)]["cra"]["approved"] == 1       # V4, though ICT rejected it
    # Fully accepted in the month the *second* authority cleared.
    assert by[(1405, 5)]["village"]["approved"] == 0
    assert by[(1405, 6)]["village"]["approved"] == 1   # V1


def test_undated_approvals_are_an_opening_balance_only(world):
    data = _progress("test_pm")
    first = data["months"][0]
    # V3 (undated) and V6 (1403, before the window) are already in the first
    # month's running total; neither is in any month's count.
    assert first["village"]["approved_cumulative"] == 2
    assert first["ict"]["approved_cumulative"] == 2
    assert first["cra"]["approved_cumulative"] == 2
    for key in STREAMS:
        in_window = sum(m[key]["approved"] for m in data["months"])
        assert data["months"][-1][key]["approved_cumulative"] == 2 + in_window


def test_window_is_twelve_months_ending_today_and_bounded(world):
    data = _progress("test_pm")
    assert len(data["months"]) == 12
    assert (data["months"][0]["shamsi_year"], data["months"][0]["shamsi_month"]) == (1404, 8)
    assert data["months"][-1]["label"] == "مهر"
    assert len(_progress("test_pm", months=99)["months"]) == 24


# ---------------------------------------------------------------------------
# 2. Reconciles with the overview
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "scope",
    [{}, {"contractor_id": "c1"}, {"contractor_id": "c2"}, {"province_id": "p1"}],
)
def test_last_running_total_equals_the_overview(client, world, scope):
    params = {k: world[v] for k, v in scope.items()}
    overview = client.get(f"{ACC}/overview", headers=world["pm"], params=params).json()
    progress = client.get(f"{ACC}/progress", headers=world["pm"], params=params)
    assert progress.status_code == 200, progress.text
    last = progress.json()["months"][-1]
    assert last["village"]["approved_cumulative"] == overview["analysis"]["villages_both_approved"]
    assert last["ict"]["approved_cumulative"] == overview["kpis"]["total_ict_approval"]
    assert last["cra"]["approved_cumulative"] == overview["kpis"]["total_cra_approval"]


def test_overview_remaining_splits_into_rejected_and_waiting(client, world):
    kpis = client.get(f"{ACC}/overview", headers=world["pm"]).json()["kpis"]
    for authority in ("ict", "cra"):
        assert (
            kpis[f"total_{authority}_rejected"] + kpis[f"total_{authority}_pending"]
            == kpis[f"total_{authority}_remained"]
        )


# ---------------------------------------------------------------------------
# 3. Plans
# ---------------------------------------------------------------------------
def test_no_plan_is_null_not_zero(world):
    by = _progress("test_pm")["by"]
    month = by[(1405, 5)]["village"]
    assert month["internal_plan"] is None
    assert month["internal_plan_cumulative"] is None
    assert month["contractor_plan"] is None
    assert month["contractor_plan_cumulative"] is None


def test_cumulative_plans_are_anchored_on_actuals(world):
    by = _progress("test_pm")["by"]
    # Internal: two approved before شهریور (V3 undated, V6), then +5, +4.
    assert by[(1405, 6)]["village"]["internal_plan"] == 5
    assert by[(1405, 6)]["village"]["internal_plan_cumulative"] == 2 + 5
    assert by[(1405, 7)]["village"]["internal_plan_cumulative"] == 2 + 5 + 4
    # Contractor: three approved before مهر (V1 joined in شهریور), then +4.
    assert by[(1405, 7)]["village"]["contractor_plan_cumulative"] == 3 + 4


def test_contractor_pip_is_the_sum_or_one_contractors(world):
    whole = _progress("test_pm")["by"][(1405, 7)]["village"]
    # C1's version 2 (3) plus C2 (1); C2's submitted شهریور PIP never counts.
    assert whole["contractor_plan"] == 4
    assert _progress("test_pm")["by"][(1405, 6)]["village"]["contractor_plan"] is None

    one = _progress("test_pm", contractor_id=world["c2"])
    assert one["by"][(1405, 7)]["village"]["contractor_plan"] == 1
    # Staff narrowed to one contractor do not see the Internal PIP.
    assert one["internal_visible"] is False
    assert one["by"][(1405, 7)]["village"]["internal_plan"] is None


# ---------------------------------------------------------------------------
# 4. A contractor
# ---------------------------------------------------------------------------
def test_contractor_never_gets_internal_and_always_gets_its_own_pip(client, world):
    for params in ({}, {"contractor_id": world["c2"]}):
        r = client.get(f"{ACC}/progress", headers=world["contractor"], params=params)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["internal_visible"] is False
        for month in body["months"]:
            for key in STREAMS:
                assert month[key]["internal_plan"] is None
                assert month[key]["internal_plan_cumulative"] is None
        assert body["months"][-1]["village"]["contractor_plan"] == 3  # C1's own
        # Its own villages only: V1 and V2 were ICT-approved; V6 is C2's.
        assert body["months"][-1]["ict"]["approved_cumulative"] == 2


def test_contractor_export_has_no_internal_pip_column(client, world):
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get(f"{ACC}/progress/export", headers=world["contractor"])
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Village", "ICT", "CRA"]
    headers = [c.value for c in wb["Village"][4]]
    assert "Contractor PIP" in headers
    assert not any("Internal" in (h or "") for h in headers)

    staff = load_workbook(BytesIO(client.get(f"{ACC}/progress/export", headers=world["pm"]).content))
    assert "Internal PIP" in [c.value for c in staff["ICT"][4]]


# ---------------------------------------------------------------------------
# 5. A province scope
# ---------------------------------------------------------------------------
def test_province_scope_has_no_plans(world):
    data = _progress("test_pm", province_id=world["p1"])
    assert data["plans_available"] is False
    for month in data["months"]:
        for key in STREAMS:
            for field in (
                "internal_plan", "internal_plan_cumulative",
                "contractor_plan", "contractor_plan_cumulative",
            ):
                assert month[key][field] is None
    # Still counts that province's approvals: V1, V2 and V6 are in P1.
    assert data["months"][-1]["ict"]["approved_cumulative"] == 3


# ---------------------------------------------------------------------------
# 6. The calendar
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "today, expected",
    [
        ((1405, 1, 15), (1405, 1, 15, 31)),
        ((1405, 7, 7), (1405, 7, 7, 30)),
        ((1404, 12, 29), (1404, 12, 29, 29)),
    ],
)
def test_today_and_month_lengths(world, today, expected):
    data = _progress("test_pm", today=jalali.from_shamsi_date(*today))
    t = data["today"]
    assert (t["shamsi_year"], t["shamsi_month"], t["day"], t["days_in_month"]) == expected
    assert [m["is_current"] for m in data["months"]] == [False] * 11 + [True]
    assert data["months"][-1]["days_in_month"] == expected[3]


# ---------------------------------------------------------------------------
# 7. ICT and CRA plan streams
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("stream, key", [("ICT", "ict"), ("CRA", "cra")])
def test_ict_and_cra_internal_pip_round_trips_into_the_chart(client, world, stream, key):
    year, month, _ = jalali.to_shamsi_date(jalali.tehran_today())
    r = client.put(
        f"{PIP}/internal-target",
        headers=world["pm"],
        json={"stream": stream, "year": year, "month": month, "target_count": 7},
    )
    assert r.status_code == 200, r.text
    read = client.get(
        f"{PIP}/internal-target", headers=world["pm"],
        params={"stream": stream, "year": year, "month": month},
    )
    assert read.json()["current"]["target_count"] == 7

    body = client.get(f"{ACC}/progress", headers=world["pm"]).json()
    assert body["months"][-1][key]["internal_plan"] == 7
    # The other acceptance stream's plan is untouched.
    other = "cra" if key == "ict" else "ict"
    assert body["months"][-1][other]["internal_plan"] in (None, 7)
    # Still never reaches a contractor.
    mine = client.get(f"{ACC}/progress", headers=world["contractor"]).json()
    assert mine["months"][-1][key]["internal_plan"] is None


@pytest.mark.parametrize("stream, key", [("ICT", "ict"), ("CRA", "cra")])
def test_ict_and_cra_contractor_pip_reaches_the_chart(world, stream, key):
    db = SessionLocal()
    _pip(db, world["c1"], stream, (1405, 7), 6)
    db.commit()
    db.close()
    by = _progress("test_pm")["by"]
    assert by[(1405, 7)][key]["contractor_plan"] == 6
    assert by[(1405, 7)]["village"]["contractor_plan"] == 4  # ACCEPTANCE unchanged


def test_unknown_stream_is_still_refused(client, world):
    r = client.put(
        f"{PIP}/internal-target", headers=world["pm"],
        json={"stream": "HC", "year": 1405, "month": 7, "target_count": 1},
    )
    assert r.status_code == 422


# ---------------------------------------------------------------------------
# The drill-through behind the ICT and CRA cards
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("authority", ["ICT", "CRA"])
def test_authority_site_lists_add_up_to_their_cards(client, world, authority):
    kpis = client.get(f"{ACC}/overview", headers=world["pm"]).json()["kpis"]
    key = authority.lower()
    expected = {
        "approved": kpis[f"total_{key}_approval"],
        "remaining": kpis[f"total_{key}_remained"],
        "rejected": kpis[f"total_{key}_rejected"],
        "remained": kpis[f"total_{key}_pending"],
    }
    for metric, total in expected.items():
        r = client.get(
            f"{ACC}/sites", headers=world["pm"],
            params={"metric": metric, "authority": authority},
        )
        assert r.status_code == 200, r.text
        assert r.json()["total"] == total, metric
        assert r.json()["label"].startswith(authority)


def test_an_unknown_authority_is_refused(client, world):
    r = client.get(f"{ACC}/sites", headers=world["pm"], params={"metric": "approved", "authority": "HQ"})
    assert r.status_code == 422
