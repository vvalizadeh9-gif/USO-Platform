"""The PIP scorecard: many Shamsi months of commitment against delivery.

``plan_and_delivery`` answers "how did this month go". This answers "how have
the last N months gone", and the difference is not the loop -- it is that a
month's workload is no longer what was handed over during it.

A contractor given twenty sites in one month and none the next is still
carrying whatever they have not finished, so the workload figure here is
``carried_in + newly_assigned``. That choice buys the right denominator and
costs the ability to add the column up, and both halves of that trade are what
these tests are for:

* **The ledger closes.** ``carried_in + newly - delivered - released ==
  carried_out``, and one month's ``carried_out`` is the next month's
  ``carried_in``. This is the property that makes every other figure
  checkable, and it is asserted on figures a real request produced, not on
  arithmetic done here.
* **A carried site is counted in every month it is open**, which is the whole
  point, and is also exactly why the balances must not be summed.
* **Reassignment ends a holding.** A site handed to another company stops
  being the first company's workload from that day, or the scorecard measures
  a contractor on work somebody else was given.
* **A drive test ends a holding too.** A finished site is not still in hand.
* **A contractor sees one company: their own.** No other name, target or
  figure anywhere in the payload, including through a site they used to hold.
* **No plan is not a zero plan.** Achievement stays None.

Run with:  cd backend && pytest tests/test_pip_scorecard.py -q
"""
import os
import sys

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_pip_scorecard_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from datetime import datetime, timezone  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402

from app.core import jalali  # noqa: E402
from app.core.database import SessionLocal  # noqa: E402
from tests.conftest import create_schema, login_form  # noqa: E402

PIP = "/api/v1/pip"
SCORECARD = f"{PIP}/scorecard"
ONAIR = "راه_اندازی_دائم"


@pytest.fixture(scope="module")
def client():
    if os.path.exists("/tmp/uep_pip_scorecard_pytest.db"):
        os.remove("/tmp/uep_pip_scorecard_pytest.db")
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        yield c


def _login(client, username="admin", password="Admin@12345"):
    r = client.post("/api/v1/auth/login", data=login_form(client, username, password))
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _role_id(client, headers, name):
    roles = client.get("/api/v1/reference/roles", headers=headers).json()
    return next(r["id"] for r in roles if r["name"] == name)


def _make_contractor(name):
    from app.models.reference import Contractor

    db = SessionLocal()
    contractor = Contractor(name=name, type="drive_test", active=True)
    db.add(contractor)
    db.commit()
    contractor_id = contractor.id
    db.close()
    return contractor_id


def _make_user(client, admin_h, username, role_name, contractor_id=None):
    password = "Test-Fixture-Passphrase"
    body = {
        "username": username,
        "password": password,
        "first_name": "Test",
        "family_name": username,
        "role_id": _role_id(client, admin_h, role_name),
        "sees_all_provinces": True,
        "province_ids": [],
    }
    if contractor_id is not None:
        body["contractor_id"] = contractor_id
    r = client.post("/api/v1/admin/users", headers=admin_h, json=body)
    assert r.status_code == 201, r.text
    return _login(client, username, password)


@pytest.fixture(scope="module")
def actors(client):
    admin_h = _login(client)
    alfa = _make_contractor("Alfa Drive Tests")
    beta = _make_contractor("Beta Surveys")

    db = SessionLocal()
    from app.models.reference import Province
    from app.models.workitem import Site

    province = Province(name="ScorecardProvince")
    db.add(province)
    db.flush()
    site = Site(site_code="SC-1", province_id=province.id)
    db.add(site)
    db.commit()
    site_id = site.id
    db.close()

    return {
        "admin": admin_h,
        "pm": _make_user(client, admin_h, "sc_pm", "PM"),
        "alfa": _make_user(client, admin_h, "sc_alfa", "Contractor", alfa),
        "beta": _make_user(client, admin_h, "sc_beta", "Contractor", beta),
        "alfa_id": alfa,
        "beta_id": beta,
        "site_id": site_id,
    }


_YEAR, _MONTH = jalali.current_shamsi_period()


def _shift(by):
    index = (_YEAR * 12 + (_MONTH - 1)) + by
    return index // 12, index % 12 + 1


def _day_in(year, month, day=10):
    return jalali.from_shamsi_date(year, month, day)


def _at(year, month, day=12):
    return datetime.combine(
        _day_in(year, month, day), datetime.min.time(), tzinfo=timezone.utc
    )


def _site(tag, *, done_in=None):
    """One on-air work item, optionally drive-test done in a Shamsi month."""
    from app.models.workitem import WorkItem

    db = SessionLocal()
    item = WorkItem(
        site_id=_ACTORS["site_id"],
        site_type=tag,
        last_stage=ONAIR,
        current_stage="New",
        dt_status="Done" if done_in else "Pending",
        dt_date_gregorian=_day_in(*done_in) if done_in else None,
    )
    db.add(item)
    db.commit()
    item_id = item.id
    db.close()
    return item_id


def _assign(work_item_id, contractor_id, year, month, day=12):
    from app.models.workitem import Assignment

    db = SessionLocal()
    db.add(
        Assignment(
            work_item_id=work_item_id,
            assignment_type="official",
            contractor_id=contractor_id,
            assigned_at=_at(year, month, day),
            is_active=True,
        )
    )
    db.commit()
    db.close()


def _approve(client, actors, key, year, month, count):
    r = client.post(
        f"{PIP}/my",
        headers=actors[key],
        json={"year": year, "month": month, "committed_count": count, "submit": True},
    )
    assert r.status_code == 200, r.text
    r = client.post(f"{PIP}/{r.json()['id']}/approve", headers=actors["pm"])
    assert r.status_code == 200, r.text


def _ask(client, headers, months=12):
    r = client.get(SCORECARD, headers=headers, params={"months": months})
    assert r.status_code == 200, r.text
    return r.json()


def _month(body, year, month):
    return next(
        m
        for m in body["months"]
        if m["shamsi_year"] == year and m["shamsi_month"] == month
    )


def _row(month, name):
    return next((r for r in month["rows"] if r["name"] == name), None)


_ACTORS: dict = {}


@pytest.fixture(scope="module", autouse=True)
def _seed(client, actors):
    """One site, carried across three months, and one that changes hands.

    Deliberately small. The ledger is a claim about how a *single* site moves
    through months, and a fixture with fifty sites in it would let a broken
    rule hide inside a total that happened to come out right.
    """
    _ACTORS.update(actors)

    y0, m0 = _shift(-4)   # assigned here, still open at the end
    y1, m1 = _shift(-3)   # carried, nothing happens
    y2, m2 = _shift(-2)   # finished here

    carried = _site("carried", done_in=(y2, m2))
    _assign(carried, actors["alfa_id"], y0, m0)

    # A second site Alfa is handed and then loses to Beta a month later.
    handed_on = _site("handed-on")
    _assign(handed_on, actors["alfa_id"], y0, m0)
    _assign(handed_on, actors["beta_id"], y1, m1)

    _approve(client, actors, "alfa", y0, m0, 2)
    _approve(client, actors, "alfa", y1, m1, 1)
    _approve(client, actors, "alfa", y2, m2, 1)
    return {"y0": (y0, m0), "y1": (y1, m1), "y2": (y2, m2)}


# ---------------------------------------------------------------------------
# 1. The ledger closes — the property everything else rests on
# ---------------------------------------------------------------------------
def test_each_month_ledger_closes(client, actors):
    body = _ask(client, actors["pm"])
    for m in body["months"]:
        assert (
            m["carried_in"] + m["newly_assigned"] - m["delivered"] - m["released"]
            == m["carried_out"]
        ), f"{m['shamsi_month_name']} {m['shamsi_year']} does not balance: {m}"


def test_carried_out_is_the_next_months_carried_in(client, actors):
    body = _ask(client, actors["pm"])
    # strict=False is the point, not an oversight: the second sequence is the
    # first shifted by one, so it is deliberately a month shorter.
    for earlier, later in zip(body["months"], body["months"][1:], strict=False):
        assert earlier["carried_out"] == later["carried_in"]


def test_contractor_rows_sum_to_their_month(client, actors):
    body = _ask(client, actors["pm"])
    for m in body["months"]:
        for field in (
            "pip",
            "carried_in",
            "newly_assigned",
            "available",
            "delivered",
            "released",
            "carried_out",
        ):
            total = sum(r[field] or 0 for r in m["rows"])
            assert total == m[field], f"{field} in {m['shamsi_month_name']}"


# ---------------------------------------------------------------------------
# 2. A carried site is workload in every month it is open
# ---------------------------------------------------------------------------
def test_a_carried_site_is_available_in_every_month_it_stays_open(client, actors, _seed):
    body = _ask(client, actors["pm"])
    y0, m0 = _seed["y0"]
    y1, m1 = _seed["y1"]
    y2, m2 = _seed["y2"]

    first = _row(_month(body, y0, m0), "Alfa Drive Tests")
    middle = _row(_month(body, y1, m1), "Alfa Drive Tests")
    last = _row(_month(body, y2, m2), "Alfa Drive Tests")

    # Handed over in the first month: new work, not carried in.
    assert first["carried_in"] == 0
    assert first["newly_assigned"] == 2
    assert first["available"] == 2

    # The middle month is the one the old rule got wrong: nothing was handed
    # over, but the site is still theirs to work on.
    assert middle["newly_assigned"] == 0
    assert middle["carried_in"] == 2
    assert middle["available"] == 1 + 1  # carried, plus handed-on until Beta takes it
    assert middle["delivered"] == 0

    # Finished in the last month, so it leaves the balance.
    assert last["delivered"] == 1
    assert last["carried_out"] == 0


def test_available_is_not_summable_and_the_payload_says_so(client, actors):
    body = _ask(client, actors["pm"])
    assert set(body["balances"]) == {"carried_in", "available", "carried_out"}
    for field in body["balances"]:
        assert field not in body["summable"]
    assert "newly_assigned" in body["summable"]
    assert "delivered" in body["summable"]


# ---------------------------------------------------------------------------
# 3. A holding ends when the site is finished, or handed to somebody else
# ---------------------------------------------------------------------------
def test_reassignment_moves_the_workload_to_the_new_contractor(client, actors, _seed):
    body = _ask(client, actors["pm"])
    y1, m1 = _seed["y1"]
    y2, m2 = _seed["y2"]

    month_of_handover = _month(body, y1, m1)
    alfa = _row(month_of_handover, "Alfa Drive Tests")
    beta = _row(month_of_handover, "Beta Surveys")

    # Alfa loses it in this month; it shows as released, not delivered.
    assert alfa["released"] == 1
    assert alfa["delivered"] == 0
    assert beta["newly_assigned"] == 1

    # And from the next month it is Beta's workload, not Alfa's.
    after = _month(body, y2, m2)
    assert _row(after, "Beta Surveys")["carried_in"] == 1
    assert _row(after, "Alfa Drive Tests")["carried_in"] == 1  # only the carried site


def test_a_finished_site_stops_being_workload(client, actors, _seed):
    body = _ask(client, actors["pm"])
    y2, m2 = _seed["y2"]
    later = [
        m
        for m in body["months"]
        if (m["shamsi_year"], m["shamsi_month"]) > (y2, m2)
    ]
    for m in later:
        alfa = _row(m, "Alfa Drive Tests")
        if alfa is not None:
            assert alfa["carried_in"] == 0, m["shamsi_month_name"]


# ---------------------------------------------------------------------------
# 4. What a contractor is shown
# ---------------------------------------------------------------------------
def test_a_contractor_sees_only_their_own_company(client, actors):
    body = _ask(client, actors["alfa"])
    assert body["is_contractor"] is True
    for m in body["months"]:
        names = {r["name"] for r in m["rows"]}
        assert names <= {"Alfa Drive Tests"}, names
    assert "Beta Surveys" not in str(body)


def test_a_contractor_is_not_measured_on_a_site_they_handed_on(client, actors, _seed):
    """Alfa's visible set still holds the handed-on site; their figures must not."""
    body = _ask(client, actors["alfa"])
    y2, m2 = _seed["y2"]
    mine = _row(_month(body, y2, m2), "Alfa Drive Tests")
    assert mine["carried_in"] == 1  # the carried site only


def test_the_pm_sees_every_company(client, actors):
    body = _ask(client, actors["pm"])
    assert body["is_contractor"] is False
    names = {r["name"] for m in body["months"] for r in m["rows"]}
    assert {"Alfa Drive Tests", "Beta Surveys"} <= names


# ---------------------------------------------------------------------------
# 5. No plan is not a zero plan
# ---------------------------------------------------------------------------
def test_no_approved_plan_leaves_achievement_unmeasured(client, actors, _seed):
    body = _ask(client, actors["pm"])
    y1, m1 = _seed["y1"]
    beta = _row(_month(body, y1, m1), "Beta Surveys")
    assert beta["pip"] is None
    assert beta["achievement_percent"] is None
    assert beta["coverage_percent"] is None
    # Execution is measurable without a plan: it is delivery over workload.
    assert beta["execution_percent"] is not None or beta["available"] == 0


def test_a_month_nobody_planned_does_not_divide_by_zero(client, actors):
    body = _ask(client, actors["pm"], months=12)
    empty = [m for m in body["months"] if m["pip"] == 0]
    assert empty, "expected at least one month with no approved plan"
    for m in empty:
        assert m["achievement_percent"] is None
        assert m["coverage_percent"] is None


# ---------------------------------------------------------------------------
# 6. The window itself
# ---------------------------------------------------------------------------
def test_rolling_window_ends_with_the_current_month(client, actors):
    body = _ask(client, actors["pm"], months=6)
    assert len(body["months"]) == 6
    last = body["months"][-1]
    assert (last["shamsi_year"], last["shamsi_month"]) == (_YEAR, _MONTH)
    years_months = [(m["shamsi_year"], m["shamsi_month"]) for m in body["months"]]
    assert years_months == sorted(years_months), "months must arrive oldest first"


def test_a_named_year_returns_all_twelve_of_its_months(client, actors):
    r = client.get(SCORECARD, headers=actors["pm"], params={"year": _YEAR})
    assert r.status_code == 200, r.text
    months = r.json()["months"]
    assert [m["shamsi_month"] for m in months] == list(range(1, 13))


# ---------------------------------------------------------------------------
# 7. The export carries the same figures, and the same scoping
# ---------------------------------------------------------------------------
def test_export_is_a_workbook_and_is_scoped_like_the_screen(client, actors):
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get(f"{PIP}/scorecard.xlsx", headers=actors["alfa"], params={"months": 6})
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]

    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Contractors"]
    text = "\n".join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    )
    assert "Beta Surveys" not in text, "a contractor's export named another company"


def test_export_does_not_total_the_balance_columns(client, actors):
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get(f"{PIP}/scorecard.xlsx", headers=actors["pm"], params={"months": 6})
    ws = load_workbook(BytesIO(r.content))["Summary"]
    total_row = next(
        row for row in ws.iter_rows() if row[0].value == "Total"
    )
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for name in ("Carried in", "Available", "Carried out"):
        assert total_row[headers.index(name)].value == "—", name
    assert isinstance(total_row[headers.index("Delivered")].value, int)


# ---------------------------------------------------------------------------
# 8. Revision history
# ---------------------------------------------------------------------------
def test_revisions_return_every_version_in_order(client, actors):
    year, month = _shift(-9)
    r = client.post(
        f"{PIP}/my",
        headers=actors["alfa"],
        json={"year": year, "month": month, "committed_count": 12, "submit": True},
    )
    plan_id = r.json()["id"]
    client.post(
        f"{PIP}/{plan_id}/return",
        headers=actors["pm"],
        json={"comment": "Too many for the sites you hold."},
    )
    client.post(
        f"{PIP}/my",
        headers=actors["alfa"],
        json={"year": year, "month": month, "committed_count": 8, "submit": True},
    )

    r = client.get(
        f"{PIP}/revisions",
        headers=actors["alfa"],
        params={"year": year, "month": month},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert [v["version"] for v in body["revisions"]] == sorted(
        v["version"] for v in body["revisions"]
    )
    assert body["revisions"][0]["return_comment"] == "Too many for the sites you hold."
    assert body["revisions"][-1]["is_current"] is True


def test_a_contractor_cannot_read_another_companys_revisions(client, actors):
    year, month = _shift(-9)
    r = client.get(
        f"{PIP}/revisions",
        headers=actors["alfa"],
        params={"year": year, "month": month, "contractor_id": actors["beta_id"]},
    )
    # The id is ignored rather than refused: answering "forbidden" would
    # confirm the other contractor exists.
    assert r.status_code == 200, r.text
    assert r.json()["contractor_id"] == actors["alfa_id"]


# ---------------------------------------------------------------------------
# 9. The contractor's own screen reads the same figures
#
# ``GET /pip/my`` carries what one screen needs: the month being filed for,
# where the running month stands, and the six months behind it. The figures in
# it are not computed a second time — they are the scorecard's, fetched for one
# contractor — and these tests are what holds that true. If the screen and the
# dashboard could disagree about a contractor's month, the argument that
# followed would be about which one to believe.
# ---------------------------------------------------------------------------
def _mine(client, headers, year=None, month=None):
    if year is None:
        year, month = _shift(1)
    r = client.get(PIP + "/my", headers=headers, params={"year": year, "month": month})
    assert r.status_code == 200, r.text
    return r.json()


def test_my_payload_history_runs_oldest_first_ending_with_the_running_month(
    client, actors
):
    body = _mine(client, actors["alfa"])
    history = body["history"]
    assert len(history) == 6

    periods = [(h["shamsi_year"], h["shamsi_month"]) for h in history]
    assert periods == sorted(periods), "history must arrive oldest first"
    assert periods[-1] == (_YEAR, _MONTH)

    # Exactly one month is still being worked on, and it is the last one.
    assert [h["in_progress"] for h in history] == [False] * 5 + [True]


def test_my_payload_assignment_matches_what_the_scorecard_reports(client, actors):
    """The whole reason the payload calls the scorecard instead of querying."""
    mine = _mine(client, actors["alfa"])
    scorecard = _ask(client, actors["alfa"], months=6)

    for point in mine["history"]:
        row = _row(
            _month(scorecard, point["shamsi_year"], point["shamsi_month"]),
            "Alfa Drive Tests",
        )
        expected = row["available"] if row else 0
        assert point["assignment"] == expected, point["label"]
        assert point["delivered"] == (row["delivered"] if row else 0), point["label"]
        assert point["pip"] == (row["pip"] if row else None), point["label"]

    running = mine["current_month"]
    row = _row(_month(scorecard, _YEAR, _MONTH), "Alfa Drive Tests")
    assert running["assignment"] == (row["available"] if row else 0)
    # And the split the screen shows underneath it adds up to it.
    assert running["carried_in"] + running["newly_assigned"] == running["assignment"]


def test_my_payload_reports_an_unapproved_month_as_null_not_zero(client, actors):
    """A contractor with no approved plan has not committed to nothing."""
    mine = _mine(client, actors["alfa"])
    by_period = {(h["shamsi_year"], h["shamsi_month"]): h for h in mine["history"]}

    # The seeded months were approved; the two most recent were never filed.
    assert by_period[_seeded_month(0)]["pip"] == 2
    assert by_period[(_YEAR, _MONTH)]["pip"] is None
    assert mine["current_month"]["pip"] is None


def _seeded_month(index):
    return _shift(-4 + index)


def test_my_payload_carries_the_planning_month_and_its_clock(client, actors):
    year, month = _shift(1)
    mine = _mine(client, actors["alfa"], year, month)
    planning = mine["planning"]

    assert (planning["shamsi_year"], planning["shamsi_month"]) == (year, month)
    assert planning["label"] == f"{jalali.month_name(month)} {year}"
    # Day 3 of a month that has not started yet is still ahead of us.
    assert planning["days_remaining"] > 0
    assert planning["deadline_passed"] is False
    # Nothing filed for it yet, and that is not an error.
    assert planning["status"] is None
    assert planning["committed_count"] is None

    # Pace is a share of the running month, so it is a percentage and nothing
    # decides anything from it.
    assert 0 <= mine["current_month"]["pace_pct"] <= 100


def test_my_payload_names_only_the_contractor_asking(client, actors):
    """No parameter names a company, and none can be smuggled in as one."""
    alfa = _mine(client, actors["alfa"])
    beta = _mine(client, actors["beta"])
    assert "Beta Surveys" not in str(alfa)

    year, month = _shift(1)
    r = client.get(
        PIP + "/my",
        headers=actors["alfa"],
        params={"year": year, "month": month, "contractor_id": actors["beta_id"]},
    )
    assert r.status_code == 200, r.text
    # The extra parameter changes nothing: the contractor is read off the
    # account, so the payload is Alfa's either way.
    assert r.json()["history"] == alfa["history"]
    assert beta["history"] != alfa["history"]


def test_a_staff_account_has_no_my_payload_to_read(client, actors):
    year, month = _shift(1)
    r = client.get(
        PIP + "/my", headers=actors["pm"], params={"year": year, "month": month}
    )
    assert r.status_code == 403, r.text


def test_pip_and_delivered_agree_with_the_drive_test_dashboard(client, actors):
    """The contractor's screen and the dashboard, on the same month.

    Two of the three figures are the same computation reached by two routes,
    and this is what says so. ``assignment`` is deliberately absent from this
    comparison: the dashboard's ``assigned`` is a *flow* — what was handed over
    during the month — and the screen's Assignment is a *stock*, what was held
    during it. They are different quantities with different names, not one
    quantity disagreeing with itself, and the screen shows the flow underneath
    the stock as "+ N new" so both are on the page.
    """
    year, month = _shift(-2)
    mine = _mine(client, actors["alfa"], *_shift(1))
    point = next(
        h
        for h in mine["history"]
        if (h["shamsi_year"], h["shamsi_month"]) == (year, month)
    )

    r = client.get(
        "/api/v1/drive-test/plan-delivery",
        headers=actors["alfa"],
        params={"year": year, "month": month},
    )
    assert r.status_code == 200, r.text
    dashboard = r.json()
    row = next(x for x in dashboard["rows"] if x["contractor_id"] == actors["alfa_id"])

    assert point["pip"] == row["pip"]
    assert point["delivered"] == row["actual"]


# ---------------------------------------------------------------------------
# 10. The PM's queue reads the same three figures
#
# The queue is about two months at once: the plan fields are the month being
# decided, and Assignment/PIP/Delivered are the month now running, which is
# what makes a proposed number credible or not. Both come from the same
# scorecard the contractor's own screen reads — these tests are what stops the
# two screens drifting into two answers about one company.
# ---------------------------------------------------------------------------
def _queue(client, headers, year=None, month=None):
    if year is None:
        year, month = _shift(1)
    r = client.get(PIP + "/queue", headers=headers, params={"year": year, "month": month})
    assert r.status_code == 200, r.text
    return r.json()


def _queue_row(body, name):
    return next(r for r in body["rows"] if r["contractor_name"] == name)


def test_queue_rows_carry_the_running_month_figures(client, actors):
    body = _queue(client, actors["pm"])
    scorecard = _ask(client, actors["pm"], months=1)
    running = _month(scorecard, _YEAR, _MONTH)

    for row in body["rows"]:
        mine = next(
            (r for r in running["rows"] if r["contractor_id"] == row["contractor_id"]),
            None,
        )
        assert row["assignment"] == (mine["available"] if mine else 0), row["contractor_name"]
        assert row["delivered"] == (mine["delivered"] if mine else 0), row["contractor_name"]
        assert row["pip"] == (mine["pip"] if mine else None), row["contractor_name"]


def test_queue_totals_are_the_programme_not_a_sum_of_what_was_shown(client, actors):
    body = _queue(client, actors["pm"])
    scorecard = _ask(client, actors["pm"], months=1)
    running = _month(scorecard, _YEAR, _MONTH)
    standing = body["current_month"]

    assert standing["assignment"] == running["available"]
    assert standing["delivered"] == running["delivered"]
    assert standing["carried_in"] + standing["newly_assigned"] == standing["assignment"]
    # Zero approved is nobody approved, not a programme that committed to none.
    assert standing["pip"] == (running["pip"] or None)
    assert 0 <= standing["pace_pct"] <= 100


def test_the_queue_names_the_month_being_decided_not_the_one_running(client, actors):
    year, month = _shift(1)
    body = _queue(client, actors["pm"], year, month)

    assert (body["shamsi_year"], body["shamsi_month"]) == (year, month)
    assert body["label"] == f"{jalali.month_name(month)} {year}"
    assert body["days_remaining"] > 0
    # ...while the standing beside it is this month.
    assert (
        body["current_month"]["shamsi_year"],
        body["current_month"]["shamsi_month"],
    ) == (_YEAR, _MONTH)


def test_a_contractor_cannot_read_the_queue_at_all(client, actors):
    year, month = _shift(1)
    r = client.get(
        PIP + "/queue", headers=actors["alfa"], params={"year": year, "month": month}
    )
    # The queue is every competitor's commitments. Not filtered — refused.
    assert r.status_code == 403, r.text


def test_admin_reads_the_queue_and_cannot_decide_a_row(client, actors):
    """Separation of duties: Admin is a systems role, not an operational one."""
    body = _queue(client, actors["admin"])
    assert body["rows"], "admin should still read the queue"

    year, month = _shift(2)
    plan_id = client.post(
        f"{PIP}/my",
        headers=actors["alfa"],
        json={"year": year, "month": month, "committed_count": 7, "submit": True},
    ).json()["id"]
    refused = client.post(f"{PIP}/{plan_id}/approve", headers=actors["admin"])
    assert refused.status_code == 403, refused.text
