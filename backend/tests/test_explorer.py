"""Explorer's site record (GET /explorer/sites/{id}) and its journey rules.

Run with:  cd backend && pytest tests/test_explorer.py -q
"""
import os
import sys
from datetime import date, datetime, timezone

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_explorer_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON, event  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402

from app.services import explorer_journey as J  # noqa: E402
from tests.conftest import create_schema, sample_cpm_path  # noqa: E402

URL = "/api/v1/explorer/sites/{}"
TODAY = date(2026, 10, 14)
ON_AIR = "راه_اندازی_دائم"
HADAF = "هدف"


def _at(d: date, hour: int = 8) -> datetime:
    return datetime(d.year, d.month, d.day, hour, tzinfo=timezone.utc)


# --------------------------------------------------------------------------
# The world
# --------------------------------------------------------------------------
@pytest.fixture(scope="module")
def world():
    if os.path.exists("/tmp/uep_explorer_pytest.db"):
        os.remove("/tmp/uep_explorer_pytest.db")
    create_schema()
    from app.core.bootstrap import init_db

    init_db()
    from app.core import user_status
    from app.core.database import SessionLocal
    from app.core.deps import CONTRACTOR, COORDINATOR
    from app.core.security import create_access_token
    from app.main import app
    from app.models.acceptance import Acceptance
    from app.models.acceptance_workflow import (
        AcceptanceAuthorityRequest,
        AcceptanceSubmission,
    )
    from app.models.health_check import HcAssignment, HcTask
    from app.models.reference import Contractor, Province, Role, User
    from app.models.workitem import Assignment, Site, Village, WorkItem

    db = SessionLocal()
    home, away = db.query(Province).order_by(Province.id).limit(2).all()
    roles = {r.name: r for r in db.query(Role).all()}
    dt_co = Contractor(name="Drive Co", type="drive_test")
    builder = Contractor(name="Builder Co", type="site")
    other_co = Contractor(name="Other Co", type="drive_test")
    db.add_all([dt_co, builder, other_co])
    db.flush()

    def work_item(code, province, **kw):
        site = Site(site_code=code, province_id=province.id)
        db.add(site)
        db.flush()
        wi = WorkItem(site_id=site.id, site_type="Greenfield", **kw)
        db.add(wi)
        db.flush()
        return wi

    def village(wi, code, *, ict="NotFiled", cra="NotFiled", flag=HADAF, deleted=False):
        v = Village(
            work_item_id=wi.id, village_code=code, village_name=f"Village {code}",
            target_classification=flag, ict_status=ict, cra_status=cra,
            deleted_at=_at(TODAY) if deleted else None,
        )
        db.add(v)
        db.flush()
        return v

    def accept(v, tech, *, ict=("Pending", None), cra=("Pending", None)):
        db.add(Acceptance(
            village_id=v.id, technology=tech,
            ict_status=ict[0], ict_date=ict[1], cra_status=cra[0], cra_date=cra[1],
        ))

    def submission(v, authority, round_no, status, submitted, reviewed=None):
        db.add(AcceptanceSubmission(
            village_id=v.id, authority=authority, round_no=round_no,
            letter_number=f"L-{v.id}-{authority}-{round_no}", letter_date=submitted,
            source="Contractor", review_status=status,
            submitted_at=_at(submitted), reviewed_at=_at(reviewed) if reviewed else None,
        ))

    # --- W4318: the full journey, acceptance in progress ----------------
    main = work_item(
        "W4318", home, last_stage=ON_AIR, launch_date_gregorian=date(2026, 9, 2),
        requested_technology="3G4G", dt_status="Done", dt_date_gregorian=date(2026, 9, 30),
        dt_sc_contractor_id=dt_co.id, site_contractor_id=builder.id,
        current_stage="DT Done",
    )
    hca = HcAssignment(code="HC-EXP-1", contractor_id=dt_co.id, assigned_at=_at(date(2026, 9, 5)))
    db.add(hca)
    db.flush()
    db.add(HcTask(hc_assignment_id=hca.id, work_item_id=main.id, overall_result="Ready",
                  completed_at=_at(date(2026, 9, 9))))
    db.add(Assignment(work_item_id=main.id, assignment_type="official", contractor_id=dt_co.id,
                      assigned_at=_at(date(2026, 9, 16)), is_active=False))

    v1 = village(main, "V10001", ict="Approved", cra="Approved")
    for tech, i, c in (("3G", 3, 5), ("4G", 4, 6)):
        accept(v1, tech, ict=("Approved", date(2026, 10, i)), cra=("Approved", date(2026, 10, c)))
    db.add(AcceptanceAuthorityRequest(village_id=v1.id, authority="ICT",
                                      sent_at=_at(date(2026, 9, 30))))

    v2 = village(main, "V10002", ict="Rejected", cra="Pending")
    accept(v2, "3G", ict=("Rejected", date(2026, 10, 2)))
    submission(v2, "ICT", 1, "Validated", date(2026, 10, 1), date(2026, 10, 2))
    submission(v2, "CRA", 1, "Pending", date(2026, 10, 7))

    v3 = village(main, "V10003", ict="Pending", cra="NotFiled")
    submission(v3, "ICT", 1, "Validated", date(2026, 10, 1), date(2026, 10, 3))
    submission(v3, "ICT", 2, "Pending", date(2026, 10, 10))

    village(main, "V10004", flag="اقماری")          # not هدف
    village(main, "V10005", deleted=True)           # removed

    # --- W0001: CPM data with no dates at all --------------------------
    undated = work_item("W0001", home, last_stage=ON_AIR, dt_status="Done",
                        current_stage="New")

    # --- W0002: fully approved ----------------------------------------
    done = work_item("W0002", home, last_stage=ON_AIR, launch_date_gregorian=date(2026, 8, 1),
                     requested_technology="4G", dt_status="Done",
                     dt_date_gregorian=date(2026, 8, 20), current_stage="DT Done")
    vd = village(done, "V20001", ict="Approved", cra="Approved")
    accept(vd, "4G", ict=("Approved", date(2026, 9, 1)), cra=("Approved", date(2026, 9, 12)))

    # --- Elsewhere: another province, another contractor ---------------
    hidden = work_item("W9999", away, dt_sc_contractor_id=other_co.id)

    # --- Query-count pair: 1 village vs 20 villages ---------------------
    def busy_site(code, n):
        wi = work_item(code, home, last_stage=ON_AIR, requested_technology="3G4G",
                       dt_status="Done", dt_date_gregorian=date(2026, 9, 1))
        for i in range(n):
            v = village(wi, f"{code}-{i:02d}", ict="Pending", cra="Approved")
            accept(v, "3G", cra=("Approved", date(2026, 9, 20)))
            accept(v, "4G", cra=("Approved", date(2026, 9, 21)))
            submission(v, "ICT", 1, "Pending", date(2026, 10, 1))
            db.add(AcceptanceAuthorityRequest(village_id=v.id, authority="CRA",
                                              sent_at=_at(date(2026, 9, 10))))
        return wi

    one, twenty = busy_site("WQ001", 1), busy_site("WQ020", 20)

    def user(username, role, **kw):
        u = User(username=username, password_hash="x", first_name=username, family_name="T",
                 role_id=roles[role].id, status=user_status.ACTIVE, **kw)
        db.add(u)
        db.flush()
        return u

    coord = user("x_coord", COORDINATOR, sees_all_provinces=False)
    coord.provinces = [home]
    sc = user("x_sc", CONTRACTOR, sees_all_provinces=False, contractor_id=dt_co.id)
    db.commit()

    def headers(u):
        return {"Authorization": f"Bearer {create_access_token(u.id, {'ver': u.token_version})}"}

    out = {
        "client": TestClient(app),
        "coord": headers(coord), "sc": headers(sc),
        "coord_id": coord.id, "sc_id": sc.id,
        "main": main.id, "undated": undated.id, "done": done.id, "hidden": hidden.id,
        "one": one.id, "twenty": twenty.id,
    }
    db.close()
    return out


def _record(world, key, who="coord"):
    """The service's answer on the fixed TODAY, read as the given user."""
    from app.core.database import SessionLocal
    from app.models.reference import User
    from app.services.explorer import site_record

    db = SessionLocal()
    try:
        user = db.get(User, world[f"{who}_id"])
        return site_record(db, user, world[key], today=TODAY)
    finally:
        db.close()


def _villages(record):
    return {v["code"]: v for v in record["villages"]}


# --------------------------------------------------------------------------
# The endpoint
# --------------------------------------------------------------------------
def test_response_shape(world):
    response = world["client"].get(URL.format(world["main"]), headers=world["coord"])
    assert response.status_code == 200, response.text
    body = response.json()

    assert set(body) == {
        "id", "site_code", "status", "general", "journey", "acceptance_started",
        "villages", "counts", "workflow",
    }
    assert body["site_code"] == "W4318"
    assert set(body["general"]) == {
        "on_air_date", "dt_date", "site_contractor", "dt_contractor",
        "requested_technology", "province", "hadaf_villages",
    }
    assert [s["key"] for s in body["journey"]["steps"]] == list(J.STEP_ORDER)
    assert set(body["journey"]["aging"]) == {
        "total_days", "current_step_days", "longest_step", "oldest_pending",
    }
    # "from" travels under its own name, not the Python-safe one.
    assert set(body["journey"]["aging"]["longest_step"]) == {"from", "to", "days"}
    assert set(body["villages"][0]["ict"]) == {
        "status", "phase", "filed_date", "decided_date", "days", "round", "filed", "refiled",
    }
    assert body["general"]["dt_date"] == "2026-09-30"


def test_a_site_outside_the_province_grant_is_404(world):
    response = world["client"].get(URL.format(world["hidden"]), headers=world["coord"])
    assert response.status_code == 404


def test_a_contractor_sees_only_its_own_sites(world):
    client = world["client"]
    assert client.get(URL.format(world["main"]), headers=world["sc"]).status_code == 200
    assert client.get(URL.format(world["hidden"]), headers=world["sc"]).status_code == 404
    assert client.get(URL.format(world["done"]), headers=world["sc"]).status_code == 404


def test_unknown_site_is_404(world):
    assert world["client"].get(URL.format(999999), headers=world["coord"]).status_code == 404


# --------------------------------------------------------------------------
# General information
# --------------------------------------------------------------------------
def test_general_information(world):
    general = _record(world, "main")["general"]
    assert general["on_air_date"] == date(2026, 9, 2)
    assert general["dt_date"] == date(2026, 9, 30)
    assert general["site_contractor"] == "Builder Co"
    assert general["dt_contractor"] == "Drive Co"
    assert general["requested_technology"] == ["3G", "4G"]
    assert general["province"]


def test_hadaf_count_is_the_acceptance_universe_rule(world):
    from app.core.database import SessionLocal
    from app.models.reference import User
    from app.services import acceptance_universe as U

    db = SessionLocal()
    try:
        user = db.get(User, world["coord_id"])
        row = next(w for w in U.load(db, user) if w.id == world["main"])
        universe = [v for v in row.villages if U.in_dt_done_universe(row, v)]
    finally:
        db.close()
    record = _record(world, "main")
    assert record["general"]["hadaf_villages"] == len(universe) == 3
    assert sorted(_villages(record)) == ["V10001", "V10002", "V10003"]


def test_on_air_date_falls_back_to_the_shamsi_cell():
    from app.models.workitem import WorkItem
    from app.services.explorer import on_air_date

    assert on_air_date(WorkItem(launch_date_shamsi="1405/07/14")) == date(2026, 10, 6)
    assert on_air_date(WorkItem(launch_date_shamsi="not a date")) is None
    assert on_air_date(WorkItem()) is None


# --------------------------------------------------------------------------
# Village statuses
# --------------------------------------------------------------------------
def test_approved_carries_decision_date_and_days_from_the_request(world):
    ict = _villages(_record(world, "main"))["V10001"]["ict"]
    assert ict["status"] == "approved"
    assert ict["decided_date"] == date(2026, 10, 4)     # the last tech to clear
    assert ict["filed_date"] == date(2026, 9, 30)
    assert ict["days"] == 4


def test_rejected_and_not_refiled(world):
    ict = _villages(_record(world, "main"))["V10002"]["ict"]
    assert (ict["status"], ict["phase"]) == ("rejected", "rejected")
    assert ict["decided_date"] == date(2026, 10, 2)
    assert ict["refiled"] is False


def test_filed_answer_waits_with_the_coordinator(world):
    cra = _villages(_record(world, "main"))["V10002"]["cra"]
    assert (cra["status"], cra["phase"]) == ("pending", "with_coordinator")
    assert cra["filed_date"] == date(2026, 10, 7)
    assert cra["days"] == 7
    assert cra["refiled"] is False


def test_refiled_round_two_is_pending_and_flagged(world):
    ict = _villages(_record(world, "main"))["V10003"]["ict"]
    assert (ict["status"], ict["round"], ict["refiled"]) == ("pending", 2, True)
    assert ict["filed_date"] == date(2026, 10, 10)
    assert ict["days"] == 4


def test_not_filed_waits_from_the_drive_test(world):
    cra = _villages(_record(world, "main"))["V10003"]["cra"]
    assert (cra["status"], cra["phase"], cra["filed"]) == ("pending", "not_filed", False)
    assert cra["days"] == 14


def test_counts_are_the_sum_of_the_sides(world):
    record = _record(world, "main")
    assert record["counts"] == {
        "ict": {"approved": 1, "pending": 1, "rejected": 1},
        "cra": {"approved": 1, "pending": 2, "rejected": 0},
    }
    for authority in ("ict", "cra"):
        assert sum(record["counts"][authority].values()) == len(record["villages"])


# --------------------------------------------------------------------------
# Journey and aging
# --------------------------------------------------------------------------
def test_steps_in_process_order_with_gaps(world):
    record = _record(world, "main")
    steps = record["journey"]["steps"]
    assert [(s["key"], s["state"], s["date"], s["gap_days"]) for s in steps] == [
        ("on_air", "done", date(2026, 9, 2), None),
        ("health_check", "done", date(2026, 9, 9), 7),
        ("assigned", "done", date(2026, 9, 16), 7),
        ("dt_done", "done", date(2026, 9, 30), 14),
        ("acceptance", "in_progress", None, 14),
    ]
    assert record["status"] == {"step": "acceptance", "state": "in_progress"}


def test_aging(world):
    aging = _record(world, "main")["journey"]["aging"]
    assert aging["total_days"] == 42
    assert aging["current_step_days"] == 14
    assert aging["longest_step"] == {"from": "assigned", "to": "dt_done", "days": 14}
    assert aging["oldest_pending"] == {"authority": "CRA", "days": 14, "villages": 2}


def test_no_dates_means_no_numbers(world):
    record = _record(world, "undated")
    steps = record["journey"]["steps"]
    assert [s["state"] for s in steps[:4]] == ["done"] * 4
    assert all(s["date"] is None and s["gap_days"] is None for s in steps)
    assert record["journey"]["aging"] == {
        "total_days": None, "current_step_days": None,
        "longest_step": None, "oldest_pending": None,
    }
    assert record["villages"] == [] and record["general"]["hadaf_villages"] == 0


def test_acceptance_done_dates_the_last_approval(world):
    record = _record(world, "done")
    assert record["status"] == {"step": "acceptance", "state": "done"}
    last = record["journey"]["steps"][-1]
    assert (last["state"], last["date"], last["gap_days"]) == ("done", date(2026, 9, 12), 23)
    assert record["journey"]["aging"]["total_days"] == 42     # 1 Aug -> 12 Sep
    assert record["journey"]["aging"]["oldest_pending"] is None


# --------------------------------------------------------------------------
# Performance
# --------------------------------------------------------------------------
def test_query_count_does_not_grow_with_villages(world):
    from app.core.database import engine

    def statements(key):
        seen = []

        def count(*_args, **_kw):
            seen.append(1)

        event.listen(engine, "before_cursor_execute", count)
        try:
            record = _record(world, key)
        finally:
            event.remove(engine, "before_cursor_execute", count)
        return len(seen), record

    small, one = statements("one")
    large, twenty = statements("twenty")
    assert len(one["villages"]) == 1 and len(twenty["villages"]) == 20
    assert small == large


# --------------------------------------------------------------------------
# Pure rules
# --------------------------------------------------------------------------
def test_a_step_dated_before_the_previous_one_has_no_gap():
    steps = J.build_steps([
        J.StepFact(J.STEP_ON_AIR, True, date(2026, 9, 10)),
        J.StepFact(J.STEP_HEALTH_CHECK, True, date(2026, 9, 1)),
        J.StepFact(J.STEP_ASSIGNED, False, None),
        J.StepFact(J.STEP_DT_DONE, False, None),
        J.StepFact(J.STEP_ACCEPTANCE, False, None),
    ], TODAY)
    assert steps[1]["gap_days"] is None
    assert steps[2]["state"] == "in_progress" and steps[2]["gap_days"] == 43
    assert steps[3]["state"] == "upcoming"


def test_a_later_step_proves_an_earlier_one():
    steps = J.build_steps([
        J.StepFact(J.STEP_ON_AIR, True, None),
        J.StepFact(J.STEP_HEALTH_CHECK, False, None),
        J.StepFact(J.STEP_ASSIGNED, False, None),
        J.StepFact(J.STEP_DT_DONE, True, date(2026, 9, 1)),
        J.StepFact(J.STEP_ACCEPTANCE, False, None),
    ], TODAY)
    assert [s["state"] for s in steps] == ["done"] * 4 + ["in_progress"]


def _side(stored, *, rounds=(), requests=(), approved=None, rejected=None):
    return J.side_status(J.SideFacts(
        stored_status=stored, approved_on=approved, rejected_on=rejected,
        rounds=tuple(rounds), requests=tuple(requests), dt_date=date(2026, 9, 1),
    ), TODAY)


def test_a_request_after_a_rejection_is_refiled_with_the_authority():
    side = _side(
        "Rejected", rejected=date(2026, 10, 1),
        rounds=[J.RoundFact(1, "Validated", date(2026, 9, 25), date(2026, 10, 1))],
        requests=[J.RequestFact(date(2026, 10, 6), None)],
    )
    assert (side["status"], side["phase"], side["round"], side["refiled"]) == (
        "pending", "with_authority", 2, True,
    )
    assert side["days"] == 8


def test_an_answered_request_is_not_open():
    side = _side(
        "NotFiled",
        rounds=[J.RoundFact(1, "Withdrawn", date(2026, 10, 2), None)],
        requests=[J.RequestFact(date(2026, 10, 1), date(2026, 9, 29))],
    )
    assert side["phase"] == "with_authority"   # a withdrawn round answers nothing
    assert side["filed_date"] == date(2026, 9, 29)


def test_returned_is_pending():
    side = _side("Returned", rounds=[J.RoundFact(1, "Returned", date(2026, 10, 1), date(2026, 10, 4))])
    assert (side["status"], side["phase"], side["days"]) == ("pending", "returned", 10)


def test_approved_without_any_request_has_no_day_count():
    side = _side("Approved", approved=date(2026, 9, 20))
    assert side["decided_date"] == date(2026, 9, 20)
    assert side["filed_date"] is None and side["days"] is None


# --------------------------------------------------------------------------
# Site contractor (CPM column 31)
# --------------------------------------------------------------------------
def test_site_contractor_is_filled_from_cpm_column_31(world, tmp_path):
    src = sample_cpm_path()
    if not src:
        pytest.skip("no sample CPM workbook available")
    import openpyxl

    from app.core.database import SessionLocal
    from app.models.reference import User
    from app.models.workitem import Site, WorkItem
    from app.services import cpm_columns as C
    from app.services.cpm_import import CpmImportService

    wb = openpyxl.load_workbook(src)
    ws = wb["CPM"]
    row = 4
    code = ws.cell(row=row, column=C.COL["site_code"] + 1).value
    stype = ws.cell(row=row, column=C.COL["site_type"] + 1).value
    expected = " ".join(str(ws.cell(row=row, column=C.COL["site_contractor"] + 1).value).split())
    path = str(tmp_path / "site_contractor.xlsx")
    wb.save(path)

    db = SessionLocal()
    try:
        admin = db.query(User).filter(User.username == "admin").first()
        CpmImportService(db, user_id=admin.id).import_file(path, "site_contractor.xlsx")
        wi = (
            db.query(WorkItem).join(Site)
            .filter(Site.site_code == str(code), WorkItem.site_type == str(stype))
            .one()
        )
        assert wi.site_contractor is not None
        assert wi.site_contractor.name == expected

        # A later import refreshes it: column 31 is a master field.
        ws.cell(row=row, column=C.COL["site_contractor"] + 1, value="Renamed Builder")
        for r in range(row + 1, ws.max_row + 1):
            if (ws.cell(row=r, column=C.COL["site_code"] + 1).value == code
                    and ws.cell(row=r, column=C.COL["site_type"] + 1).value == stype):
                ws.cell(row=r, column=C.COL["site_contractor"] + 1, value="Renamed Builder")
        wb.save(path)
        CpmImportService(db, user_id=admin.id).import_file(path, "site_contractor_2.xlsx")
        db.refresh(wi)
        assert wi.site_contractor.name == "Renamed Builder"
    finally:
        db.close()


# --------------------------------------------------------------------------
# The list's search (GET /work-items?q=)
# --------------------------------------------------------------------------
def test_list_search_matches_village_id_and_name(world):
    client = world["client"]
    for term in ("V10002", "Village V10003"):
        body = client.get("/api/v1/work-items", params={"q": term}, headers=world["coord"]).json()
        assert [r["site_code"] for r in body] == ["W4318"]
    # A removed village no longer finds its site.
    body = client.get("/api/v1/work-items", params={"q": "V10005"}, headers=world["coord"]).json()
    assert body == []
