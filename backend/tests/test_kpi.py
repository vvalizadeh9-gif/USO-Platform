"""KPI & Performance: the metric definitions, and who may see them.

The dataset below is small and hand-checked, so every assertion names the
arithmetic it expects rather than comparing against whatever the code happens
to produce. It deliberately contains the three cases that are easy to get
wrong:

* a province with too few DT-done villages to compare (Zanjan),
* a site whose CPM province cell matched none of the 31, so it belongs to no
  owner (the "Unknown province" row),
* a village that was rejected and later approved, which must count once, as
  approved.

Run with:  cd backend && pytest tests/test_kpi.py -q
"""
import os
import sys

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_kpi_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from tests.conftest import create_schema, login_as_role, login_form  # noqa: E402

DB_FILE = "/tmp/uep_kpi_pytest.db"

TEHRAN = "تهران"
MAZANDARAN = "مازندران"
ARDABIL = "اردبیل"
ZANJAN = "زنجان"

CONTRACTOR_A = "DT-Alpha"
CONTRACTOR_B = "DT-Beta"

ON_AIR = "راه_اندازی_دائم"


@pytest.fixture(scope="module")
def client():
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        _seed()
        yield c


def _villages(db, work_item_id, count, *, ict_approved=0, ict_rejected=0,
              cra_approved=0, cra_rejected=0):
    """``count`` target villages on one work item, with the given verdicts."""
    from app.models.workitem import Village

    made = []
    for index in range(count):
        ict = "NotFiled"
        if index < ict_approved:
            ict = "Approved"
        elif index < ict_approved + ict_rejected:
            ict = "Rejected"
        cra = "NotFiled"
        if index < cra_approved:
            cra = "Approved"
        elif index < cra_approved + cra_rejected:
            cra = "Rejected"
        village = Village(
            work_item_id=work_item_id,
            village_code=f"WI{work_item_id}-V{index}",
            target_classification="هدف",
            ict_status=ict,
            cra_status=cra,
        )
        db.add(village)
        made.append(village)
    db.flush()
    return made


def _seed() -> None:
    """Five work items, one per situation the page has to get right.

    Country totals that fall out of it, checked by hand:

        villages          12 + 10 + 10 + 4 + 2 = 38
        DT-done villages  12 + 10 + 10 + 0 + 2 = 34
        ICT approved       9 +  5 + 10 + 0 + 2 = 26   -> 26/34 = 76.5%
    """
    from app.models.reference import Contractor, Province
    from app.models.workitem import Site, WorkItem

    db = SessionLocal()
    try:
        provinces = {
            name: db.query(Province).filter(Province.name == name).one()
            for name in (TEHRAN, MAZANDARAN, ARDABIL, ZANJAN)
        }
        alpha = Contractor(name=CONTRACTOR_A, type="drive_test")
        beta = Contractor(name=CONTRACTOR_B, type="drive_test")
        db.add_all([alpha, beta])
        db.flush()

        def site(code, province_name):
            row = Site(
                site_code=code,
                province_id=provinces[province_name].id if province_name else None,
            )
            db.add(row)
            db.flush()
            return row

        def work_item(site_row, contractor, *, dt_done: bool):
            row = WorkItem(
                site_id=site_row.id,
                site_type="A",
                requested_technology="2G",
                last_stage=ON_AIR,
                dt_status="Done" if dt_done else "Ongoing",
                dt_sc_contractor_id=contractor.id,
                current_stage="New",
            )
            db.add(row)
            db.flush()
            return row

        # Tehran: 12 villages, 9 ICT approved (one of them approved only after
        # an earlier rejection), 2 rejected, 1 never filed. CRA: 6 approved.
        tehran = work_item(site("S-TEH", TEHRAN), alpha, dt_done=True)
        made = _villages(
            db, tehran.id, 12, ict_approved=9, ict_rejected=2, cra_approved=6
        )
        _rejected_then_approved(db, made[0])

        # Mazandaran: 10 villages, 5 ICT approved.
        mazandaran = work_item(site("S-MAZ", MAZANDARAN), beta, dt_done=True)
        _villages(db, mazandaran.id, 10, ict_approved=5, cra_approved=10)

        # Ardabil: 10 villages, all ICT approved.
        ardabil = work_item(site("S-ARD", ARDABIL), alpha, dt_done=True)
        _villages(db, ardabil.id, 10, ict_approved=10, cra_approved=3)

        # Zanjan: on air but the drive test is not done, so it has no base to
        # be compared on -- the low-sample case.
        zanjan = work_item(site("S-ZAN", ZANJAN), beta, dt_done=False)
        _villages(db, zanjan.id, 4)

        # A site whose CPM province cell matched none of the 31.
        unknown = work_item(site("S-UNK", None), alpha, dt_done=True)
        _villages(db, unknown.id, 2, ict_approved=2, cra_approved=2)

        db.commit()
    finally:
        db.close()


def _rejected_then_approved(db, village) -> None:
    """Give one village a rejected first round and an approved second one.

    The verdict is then derived the way the application derives it, rather
    than being written by hand -- which is the only way this test can show
    that "final status wins" is a property of the system and not of the
    fixture.
    """
    from datetime import date, datetime, timezone

    from app.models.acceptance import Acceptance
    from app.models.acceptance_workflow import (
        AUTHORITY_ICT,
        CLAIM_APPROVED,
        CLAIM_REJECTED,
        REVIEW_VALIDATED,
        SOURCE_COORDINATOR,
        AcceptanceSubmission,
        AcceptanceSubmissionTech,
    )
    from app.services import acceptance_workflow as flow

    village.acceptances = [
        Acceptance(technology="2G", ict_status="Approved", cra_status="Pending")
    ]
    now = datetime.now(timezone.utc)
    for round_no, claim in ((1, CLAIM_REJECTED), (2, CLAIM_APPROVED)):
        submission = AcceptanceSubmission(
            village_id=village.id,
            authority=AUTHORITY_ICT,
            round_no=round_no,
            letter_number=f"L-{round_no}",
            letter_date=date.today(),
            source=SOURCE_COORDINATOR,
            review_status=REVIEW_VALIDATED,
            submitted_at=now,
        )
        submission.technologies = [
            AcceptanceSubmissionTech(technology="2G", claimed_status=claim)
        ]
        db.add(submission)
    db.flush()
    flow.recompute_authority_statuses(db, [village])
    db.flush()


# ----- Sign-in helpers ----------------------------------------------------


def _admin(client) -> dict:
    response = client.post("/api/v1/auth/login", data=login_form(client))
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _token(client, role: str) -> dict:
    return login_as_role(client, _admin(client)["Authorization"].split()[1], role)


@pytest.fixture(scope="module")
def actors(client):
    """One signed-in account per role, with the two name links in place."""
    admin_h = _admin(client)
    pm = _token(client, "PM")
    rm = _token(client, "RegionalManager")
    coordinator = _token(client, "Coordinator")

    users = client.get("/api/v1/admin/users", headers=admin_h).json()
    by_username = {u["username"]: u for u in users}

    for username, name in (
        ("test_regionalmanager", "Pirayesh"),
        ("test_coordinator", "Hossein"),
    ):
        response = client.put(
            f"/api/v1/kpi/mapping/links/{by_username[username]['id']}",
            headers=pm,
            json={"kpi_person_name": name},
        )
        assert response.status_code == 200, response.text

    contractor_h = _contractor_account(client, admin_h)
    return {
        "admin": admin_h,
        "pm": pm,
        "rm": rm,
        "coordinator": coordinator,
        "contractor": contractor_h,
    }


def _contractor_account(client, admin_h) -> dict:
    """A contractor account pointed at DT-Alpha."""
    from app.models.reference import Contractor

    db = SessionLocal()
    contractor_id = (
        db.query(Contractor).filter(Contractor.name == CONTRACTOR_A).one().id
    )
    db.close()

    roles = client.get("/api/v1/reference/roles", headers=admin_h).json()
    role_id = next(r["id"] for r in roles if r["name"] == "Contractor")
    password = "Drive-Test-Sc-Passw0rd"
    existing = client.get("/api/v1/admin/users", headers=admin_h).json()
    if not any(u["username"] == "kpi_contractor" for u in existing):
        created = client.post(
            "/api/v1/admin/users",
            headers=admin_h,
            json={
                "username": "kpi_contractor",
                "password": password,
                "first_name": "Kpi",
                "family_name": "Contractor",
                "role_id": role_id,
                "contractor_id": contractor_id,
                "sees_all_provinces": False,
                "province_ids": [],
            },
        )
        assert created.status_code == 201, created.text
    response = client.post(
        "/api/v1/auth/login", data=login_form(client, "kpi_contractor", password)
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _summary(client, headers, **params):
    response = client.get("/api/v1/kpi/summary", headers=headers, params=params)
    assert response.status_code == 200, response.text
    return response.json()


# ----- Acceptance check 2: the country average is weighted ----------------


def test_country_average_is_weighted_not_a_mean_of_provinces(client, actors):
    payload = _summary(client, actors["pm"], lens="rm", key="Pirayesh")
    country = payload["country"]

    assert country["villages"] == 38
    assert country["villages_dt_done"] == 34
    assert country["ict_approved"] == 26
    # 26 / 34 = 76.47…, the weighted figure.
    assert country["ict_approved_pct"] == 76.5

    # The unweighted mean of the four provinces that have a base would be
    # (75.0 + 50.0 + 100.0 + 100.0) / 4 = 81.25. It is not what is reported,
    # which is the whole point of weighting.
    assert country["ict_approved_pct"] != 81.3


# ----- Acceptance check 3: rejected then approved counts once, as approved -


def test_rejected_then_approved_counts_as_approved(client, actors):
    from app.models.workitem import Village

    db = SessionLocal()
    village = db.query(Village).filter(Village.village_code.like("%-V0")).order_by(
        Village.id
    ).first()
    statuses = [
        (s.round_no, s.review_status)
        for s in _submissions_for(db, village.id)
    ]
    derived = village.ict_status
    db.close()

    # It really does carry a rejected round in its history…
    assert (1, "Validated") in statuses
    assert (2, "Validated") in statuses
    # …and the status the KPI page reads is the final one.
    assert derived == "Approved"

    payload = _summary(client, actors["pm"], lens="rm", key="Allahyar")
    tehran = _row(payload, "Tehran")
    # 9 approved out of 12 villages, 12 of them DT-done: 75.0%.
    assert tehran["ict_approved"]["count"] == 9
    assert tehran["ict_approved"]["pct"] == 75.0
    assert tehran["ict_rejected"]["count"] == 2


def _submissions_for(db, village_id):
    from app.models.acceptance_workflow import AcceptanceSubmission

    return (
        db.query(AcceptanceSubmission)
        .filter(AcceptanceSubmission.village_id == village_id)
        .all()
    )


def _row(payload, province: str) -> dict:
    return next(r for r in payload["provinces"] if r["province"] == province)


# ----- Acceptance check 1: the scopes add up to the country ---------------


def test_every_regional_manager_scope_sums_to_the_country(client, actors):
    _assert_scopes_sum(client, actors["pm"], "rm")


def test_every_coordinator_scope_sums_to_the_country(client, actors):
    _assert_scopes_sum(client, actors["pm"], "coordinator")


def _assert_scopes_sum(client, pm_headers, lens: str) -> None:
    """Every owner's scope, added up, is the country minus what nobody owns.

    A site whose CPM province cell matched none of the 31 has no province and
    therefore no owner, so it cannot appear in anybody's scope. It is counted
    in the country total and shown as its own row, and the difference is
    exactly that row — which is the property worth asserting, because a
    silently dropped site would show up here as a mismatch.
    """
    options = client.get("/api/v1/kpi/lenses", headers=pm_headers).json()
    keys = options["options"][lens]
    assert keys, f"no {lens} options"

    villages = dt_done = ict_approved = work_items = 0
    for key in keys:
        payload = _summary(client, pm_headers, lens=lens, key=key)
        villages += payload["villages"]["total"]
        dt_done += payload["villages"]["dt_done"]
        ict_approved += payload["villages"]["ict_approved"]
        work_items += payload["work_items"]["total"]

    country = _summary(client, pm_headers, lens=lens, key=keys[0])["country"]
    unowned = _unknown_province_totals(client, pm_headers)

    assert villages + unowned["villages"] == country["villages"]
    assert dt_done + unowned["dt_done"] == country["villages_dt_done"]
    assert ict_approved + unowned["ict_approved"] == country["ict_approved"]
    assert work_items + unowned["work_items"] == country["work_items"]


def _unknown_province_totals(client, pm_headers) -> dict:
    """What the "Unknown province" row holds, read from the country view.

    The contractor lens is the one lens that is not geographic, so it is the
    only place a province-less site is visible without special-casing.
    """
    seen = {"villages": 0, "dt_done": 0, "ict_approved": 0, "work_items": 0}
    for contractor in (CONTRACTOR_A, CONTRACTOR_B):
        payload = _summary(client, pm_headers, lens="contractor", key=contractor)
        for row in payload["provinces"]:
            if row["province_fa"] is None:
                seen["villages"] += row["villages"]
                seen["dt_done"] += row["dt_done_villages"]
                seen["ict_approved"] += row["ict_approved"]["count"]
                seen["work_items"] += 1
    return seen


# ----- Acceptance check 4: low-sample provinces -------------------------


def test_low_sample_province_is_shown_uncoloured_and_last(client, actors):
    payload = _summary(client, actors["pm"], lens="rm", key="Pirayesh")
    names = [row["province"] for row in payload["provinces"]]

    assert "Zanjan" in names, "a low-sample province is still shown"
    assert names[-1] == "Zanjan", "low-sample provinces sort last"

    zanjan = _row(payload, "Zanjan")
    assert zanjan["low_sample"] is True
    assert zanjan["dt_done_villages"] < payload["low_sample_threshold"]
    # Nothing to compare against, so there is no difference to colour by.
    assert zanjan["ict_approved"]["pct"] is None
    assert zanjan["ict_approved"]["delta"] is None

    ardabil = _row(payload, "Ardabil")
    assert ardabil["low_sample"] is False
    assert ardabil["ict_approved"]["pct"] == 100.0
    # 100.0 against the country's 76.5.
    assert ardabil["ict_approved"]["delta"] == 23.5


def test_low_sample_never_outranks_a_compared_province(client, actors):
    """A province with a perfect rate on three villages must not head the table."""
    from app.services.kpi import LOW_SAMPLE_DT_DONE, _sort_key

    perfect_but_tiny = {"low_sample": True, "ict_approved": {"pct": 100.0}}
    ordinary = {"low_sample": False, "ict_approved": {"pct": 10.0}}
    assert _sort_key(ordinary) < _sort_key(perfect_but_tiny)
    assert LOW_SAMPLE_DT_DONE == 10


# ----- Acceptance check 5: access is enforced on the backend --------------


def test_admin_has_no_access_anywhere(client, actors):
    admin = actors["admin"]
    for url, params in (
        ("/api/v1/kpi/summary", {}),
        ("/api/v1/kpi/lenses", {}),
        ("/api/v1/kpi/contractors", {"mode": "ict"}),
        ("/api/v1/kpi/export.xlsx", {}),
        ("/api/v1/kpi/export.pdf", {}),
        ("/api/v1/kpi/mapping", {}),
    ):
        response = client.get(url, headers=admin, params=params)
        assert response.status_code == 403, f"{url} -> {response.status_code}"


def test_regional_manager_is_confined_to_their_own_provinces(client, actors):
    rm = actors["rm"]

    own = _summary(client, rm, lens="rm", key="Pirayesh")
    assert own["key"] == "Pirayesh"
    assert own["selectable"] is False
    assert sorted(r["province"] for r in own["provinces"]) == ["Ardabil", "Zanjan"]

    # Asking with no parameters gets the same scope, not everything.
    default = _summary(client, rm)
    assert default["key"] == "Pirayesh"

    # Another manager's data, and another lens, are both refused.
    assert client.get(
        "/api/v1/kpi/summary", headers=rm, params={"lens": "rm", "key": "Allahyar"}
    ).status_code == 403
    assert client.get(
        "/api/v1/kpi/summary",
        headers=rm,
        params={"lens": "coordinator", "key": "Hossein"},
    ).status_code == 403


def test_regional_manager_cannot_see_contractor_comparison(client, actors):
    response = client.get(
        "/api/v1/kpi/contractors", headers=actors["rm"], params={"mode": "ict"}
    )
    assert response.status_code == 403


def test_contractor_sees_only_their_own_dt_sc_rows(client, actors):
    payload = _summary(client, actors["contractor"])
    assert payload["lens"] == "contractor"
    assert payload["key"] == CONTRACTOR_A

    # DT-Alpha holds Tehran, Ardabil and the province-less site: 12 + 10 + 2.
    assert payload["villages"]["total"] == 24
    provinces = {row["province"] for row in payload["provinces"]}
    assert "Mazandaran" not in provinces, "that is DT-Beta's province"

    # The benchmark is still the whole country, as the brief requires.
    assert payload["country"]["villages"] == 38

    assert client.get(
        "/api/v1/kpi/summary",
        headers=actors["contractor"],
        params={"lens": "contractor", "key": CONTRACTOR_B},
    ).status_code == 403
    assert client.get(
        "/api/v1/kpi/contractors", headers=actors["contractor"], params={"mode": "ict"}
    ).status_code == 403


def test_coordinator_gets_their_own_contractor_table(client, actors):
    response = client.get(
        "/api/v1/kpi/contractors", headers=actors["coordinator"], params={"mode": "ict"}
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["key"] == "Hossein"
    # Hossein holds Ardabil among others; only DT-Alpha works there.
    assert payload["total"]["villages"] == 10
    assert payload["total"]["approved"] == 10
    assert payload["total"]["remained"] == 0

    # Another coordinator's table is refused.
    assert client.get(
        "/api/v1/kpi/contractors",
        headers=actors["coordinator"],
        params={"key": "Amir", "mode": "ict"},
    ).status_code == 403


def test_only_pm_may_read_or_change_the_mapping(client, actors):
    assert client.get("/api/v1/kpi/mapping", headers=actors["pm"]).status_code == 200
    for role in ("rm", "coordinator", "contractor", "admin"):
        assert (
            client.get("/api/v1/kpi/mapping", headers=actors[role]).status_code == 403
        ), role


# ----- The mapping itself -------------------------------------------------


def test_mapping_seeds_all_31_provinces_and_matches_the_platform(client, actors):
    from app.core.province_directory import PROVINCE_DIRECTORY
    from app.services.cpm_columns import IRAN_PROVINCES

    payload = client.get("/api/v1/kpi/mapping", headers=actors["pm"]).json()
    assert len(payload["rows"]) == 31
    assert payload["unmapped_provinces"] == []

    # Every Persian name in the directory is one the CPM import can produce.
    # If these two lists ever drift, a province silently leaves every lens.
    assert {row.fa for row in PROVINCE_DIRECTORY} == set(IRAN_PROVINCES)
    assert len({row.en for row in PROVINCE_DIRECTORY}) == 31


def test_reassignment_closes_the_old_row_and_keeps_it(client, actors):
    # Today, not a literal. Startup opens every seeded mapping row at
    # date.today(), and a reassignment may not start before the row it
    # replaces -- so a date written into this file passes on the day it is
    # written and fails every day after it.
    from datetime import date

    today = date.today().isoformat()
    pm = actors["pm"]
    rows = client.get("/api/v1/kpi/mapping", headers=pm).json()["rows"]
    qom = next(r for r in rows if r["province_en"] == "Qom")
    assert qom["regional_manager"] == "Allahyar"

    response = client.post(
        f"/api/v1/kpi/mapping/{qom['id']}/reassign",
        headers=pm,
        json={
            "cra_region": qom["cra_region"],
            "pso_coordinator": qom["pso_coordinator"],
            "regional_manager": "Rouhi",
            "effective_from": today,
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["regional_manager"] == "Rouhi"

    history = client.get(
        "/api/v1/kpi/mapping", headers=pm, params={"include_history": True}
    ).json()["history"]["قم"]
    assert len(history) == 2
    assert history[0]["regional_manager"] == "Allahyar"
    assert history[0]["effective_to"] == today
    assert history[1]["effective_to"] is None

    # Put it back so the rest of the module sees the seeded mapping.
    fresh = client.get("/api/v1/kpi/mapping", headers=pm).json()["rows"]
    restored = next(r for r in fresh if r["province_en"] == "Qom")
    client.put(
        f"/api/v1/kpi/mapping/{restored['id']}",
        headers=pm,
        json={
            "cra_region": "Central",
            "pso_coordinator": "Hossein",
            "regional_manager": "Allahyar",
        },
    )


def test_an_unlinked_account_is_told_why_it_cannot_see_the_page(client, actors):
    """A Regional Manager with no name link is refused with an explanation."""
    from app.models.reference import User

    db = SessionLocal()
    account = db.query(User).filter(User.username == "test_regionalmanager").one()
    original = account.kpi_person_name
    account.kpi_person_name = None
    db.commit()
    db.close()

    response = client.get("/api/v1/kpi/summary", headers=actors["rm"])
    assert response.status_code == 403
    assert "province mapping" in response.json()["detail"]

    db = SessionLocal()
    account = db.query(User).filter(User.username == "test_regionalmanager").one()
    account.kpi_person_name = original
    db.commit()
    db.close()


# ----- Acceptance check 6: the exports match the screen -------------------


def test_excel_export_carries_the_same_numbers_as_the_screen(client, actors):
    import io

    from openpyxl import load_workbook

    params = {"lens": "rm", "key": "Pirayesh"}
    payload = _summary(client, actors["pm"], **params)

    response = client.get("/api/v1/kpi/export.xlsx", headers=actors["pm"], params=params)
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames[:2] == ["Summary", "Provinces"]

    summary_values = [
        [cell.value for cell in row] for row in workbook["Summary"].iter_rows()
    ]
    flat = [v for row in summary_values for v in row if v is not None]
    assert "Pirayesh" in flat
    assert payload["villages"]["total"] in flat
    assert payload["villages"]["ict_approved"] in flat

    province_values = [
        [cell.value for cell in row] for row in workbook["Provinces"].iter_rows()
    ]
    ardabil = next(row for row in province_values if row[0] == "Ardabil")
    on_screen = _row(payload, "Ardabil")
    assert ardabil[2] == on_screen["villages"]
    assert ardabil[3] == on_screen["dt_done_villages"]
    # ICT approved % is the 9th column: 4 fixed + 2 per measure, ICT approved
    # being the third measure.
    assert ardabil[8] == f'{on_screen["ict_approved"]["pct"]:.1f}%'

    zanjan = next(row for row in province_values if row[0] == "Zanjan")
    assert "not compared" in zanjan


def test_pdf_export_is_produced(client, actors):
    response = client.get(
        "/api/v1/kpi/export.pdf",
        headers=actors["pm"],
        params={"lens": "rm", "key": "Pirayesh"},
    )
    assert response.status_code == 200, response.text
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 2000
    assert "attachment" in response.headers["content-disposition"]


def test_exports_obey_the_same_scope_rules(client, actors):
    """A regional manager's export cannot hold another manager's provinces."""
    response = client.get(
        "/api/v1/kpi/export.xlsx",
        headers=actors["rm"],
        params={"lens": "rm", "key": "Allahyar"},
    )
    assert response.status_code == 403


# ----- Everything carries the CPM import time -----------------------------


def test_every_response_names_the_last_cpm_import(client, actors):
    payload = _summary(client, actors["pm"], lens="rm", key="Pirayesh")
    assert "last_cpm_import" in payload

    contractors = client.get(
        "/api/v1/kpi/contractors", headers=actors["pm"], params={"mode": "ict"}
    ).json()
    assert "last_cpm_import" in contractors

    mapping = client.get("/api/v1/kpi/mapping", headers=actors["pm"]).json()
    assert "last_cpm_import" in mapping
