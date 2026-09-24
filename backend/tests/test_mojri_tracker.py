"""Mojri tracker reconciliation: the template, the reading of a cell, and who
may do either.

Three of these tests protect a rule that, if broken, would be believed:

* **nothing is written before Confirm.** A preview that quietly wrote would
  make the whole screen a lie, and the person would have approved numbers
  after the fact;
* **a cell for a technology the village never requested changes nothing.** A
  3G/4G village's 2G column is not a gap; counting it would report a shortfall
  that can never close;
* **a village that disappears from the file is never reverted.** Somebody
  filtering a row out of a spreadsheet is not evidence that a registration was
  withdrawn.

Run with:  cd backend && pytest tests/test_mojri_tracker.py -q
"""
import io
import os
import sys

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_mojri_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.models.mojri import IN_TRACKER, NEEDS_LOOK, NOT_IN_TRACKER  # noqa: E402
from app.services import mojri_tracker  # noqa: E402
from tests.conftest import create_schema, login_as_role, login_form  # noqa: E402

DB_FILE = "/tmp/uep_mojri_pytest.db"

TEHRAN = "تهران"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: Set by _seed, read by the tests: village id -> what that village is for.
IDS = {}


@pytest.fixture(scope="module")
def client():
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        _seed()
        yield c


def _seed() -> None:
    """Five villages: the cases the importer has to tell apart.

    * ``three_tech``   2G/3G/4G, ICT approved -- the ordinary case
    * ``two_tech``     3G/4G only, ICT approved -- its 2G column must be ignored
    * ``cra_only``     CRA approved, ICT not -- still eligible for the template
    * ``no_tech``      approved, but CPM recorded no requested technology
    * ``not_eligible`` neither authority approved -- must not be in the template
    """
    from app.models.reference import Province
    from app.models.workitem import Site, Village, WorkItem

    db = SessionLocal()
    try:
        province = db.query(Province).filter(Province.name == TEHRAN).one()

        def village(label, *, tech, ict, cra, site_type):
            site = Site(site_code=f"S-{label}", province_id=province.id)
            db.add(site)
            db.flush()
            work_item = WorkItem(
                site_id=site.id,
                site_type=site_type,
                requested_technology=tech,
                last_stage="راه_اندازی_دائم",
                dt_status="Done",
                current_stage="New",
            )
            db.add(work_item)
            db.flush()
            row = Village(
                work_item_id=work_item.id,
                village_code=f"V-{label}",
                village_name=f"village {label}",
                target_classification="هدف",
                ict_status=ict,
                cra_status=cra,
            )
            db.add(row)
            db.flush()
            IDS[label] = row.id

        village("three_tech", tech="2G3G4G", ict="Approved", cra="NotFiled", site_type="A")
        village("two_tech", tech="3G4G", ict="Approved", cra="NotFiled", site_type="B")
        village("cra_only", tech="2G", ict="NotFiled", cra="Approved", site_type="C")
        village("no_tech", tech=None, ict="Approved", cra="NotFiled", site_type="D")
        village("not_eligible", tech="2G", ict="NotFiled", cra="NotFiled", site_type="E")
        db.commit()
    finally:
        db.close()


# ----- Sign-in helpers ----------------------------------------------------


def _admin(client) -> dict:
    response = client.post("/api/v1/auth/login", data=login_form(client))
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture(scope="module")
def actors(client):
    admin_h = _admin(client)
    token = admin_h["Authorization"].split()[1]
    return {
        "admin": admin_h,
        "pm": login_as_role(client, token, "PM"),
        "coordinator": login_as_role(client, token, "Coordinator"),
    }


# ----- Building a filled template ----------------------------------------


def _fill(rows, *, headers=None, comments=None):
    """A filled template as .xlsx bytes.

    ``rows`` is a list of (village_id, {column: value}); ``comments`` is
    {(row_index, column): text} for the Excel-comment case.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MojriTracker"
    header = list(headers or mojri_tracker.TEMPLATE_COLUMNS)
    sheet.append(header)
    for index, (village_id, cells) in enumerate(rows, start=2):
        line = []
        for column in header:
            if column == "village_id":
                line.append(village_id)
            elif column in ("site_id", "site_type"):
                line.append("ref")
            else:
                line.append(cells.get(column, ""))
        sheet.append(line)
        for (comment_row, column), text in (comments or {}).items():
            if comment_row == index and column in header:
                cell = sheet.cell(row=index, column=header.index(column) + 1)
                cell.comment = Comment(text, "tester")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, headers, path, content, **data):
    return client.post(
        f"/api/v1/mojri/{path}",
        headers=headers,
        files={"file": ("mojri.xlsx", content, XLSX)},
        data=data,
    )


def _preview(client, headers, content):
    response = _upload(client, headers, "import/preview", content)
    assert response.status_code == 200, response.text
    return response.json()


def _statuses(village_id):
    from app.models.mojri import MojriTrackerStatus

    db = SessionLocal()
    try:
        row = (
            db.query(MojriTrackerStatus)
            .filter(MojriTrackerStatus.village_id == village_id)
            .one_or_none()
        )
        return None if row is None else (row.ict_status, row.cra_status)
    finally:
        db.close()


# ----- The template -------------------------------------------------------


def test_the_template_holds_the_villages_we_have_already_approved(client, actors):
    response = client.get("/api/v1/mojri/template.xlsx", headers=actors["admin"])
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX

    sheet = load_workbook(io.BytesIO(response.content)).active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == list(mojri_tracker.TEMPLATE_COLUMNS)

    rows = {row[1]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {
        IDS["three_tech"],
        IDS["two_tech"],
        IDS["cra_only"],
        IDS["no_tech"],
    }
    # A village neither authority has approved is not behind in anybody's
    # tracker yet, so it is not asked about.
    assert IDS["not_eligible"] not in rows


def test_the_template_marks_technologies_a_village_never_requested(client, actors):
    response = client.get("/api/v1/mojri/template.xlsx", headers=actors["admin"])
    sheet = load_workbook(io.BytesIO(response.content)).active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = {row[1]: dict(zip(header, row, strict=True)) for row in sheet.iter_rows(min_row=2, values_only=True)}

    two_tech = rows[IDS["two_tech"]]
    assert two_tech["ict_2g"] == "n/a"      # never requested
    # Blank, ready to fill. openpyxl reads an empty cell back as None.
    assert two_tech["ict_3g"] is None
    assert two_tech["ict_4g"] is None
    # The reference columns are there for the human matching against Mojri's
    # own file, which is organised by site.
    assert two_tech["site_id"] == "S-two_tech"
    assert two_tech["site_type"] == "B"


def test_the_filename_carries_the_shamsi_period():
    name = mojri_tracker.template_filename((1404, 6))
    assert name == "mojri_template_1404_06.xlsx"


# ----- Reading a cell -----------------------------------------------------


@pytest.mark.parametrize(
    "value,expected",
    [
        ("yes", IN_TRACKER),
        ("Approved", IN_TRACKER),
        ("تایید", IN_TRACKER),
        ("✓", IN_TRACKER),
        (1, IN_TRACKER),
        ("", NOT_IN_TRACKER),
        ("   ", NOT_IN_TRACKER),
        ("?", NEEDS_LOOK),
        ("check", NEEDS_LOOK),
        ("yes - but see the letter", NEEDS_LOOK),
        ("registered 1404/05", NEEDS_LOOK),
        # A recognised rejection is NOT read as "not in the tracker": in a
        # registration column, refused and not-yet-done are different facts.
        ("no", NEEDS_LOOK),
        ("رد", NEEDS_LOOK),
    ],
)
def test_one_cell_routes_where_it_should(client, actors, value, expected):
    content = _fill(
        [(IDS["three_tech"], {"ict_2g": value, "ict_3g": value, "ict_4g": value})]
    )
    payload = _preview(client, actors["pm"], content)
    counts = payload["authorities"]["ict"]
    assert counts[expected] == 1, payload


def test_a_cell_carrying_a_comment_needs_a_look(client, actors):
    """The value reads as a clean yes; the comment on it does not.

    A note somebody attached to a cell is the most common way a real file says
    "this one is complicated", and it is invisible to anything that reads only
    the value.
    """
    content = _fill(
        [(IDS["three_tech"], {"ict_2g": "yes", "ict_3g": "yes", "ict_4g": "yes"})],
        comments={(2, "ict_3g"): "waiting on the letter"},
    )
    payload = _preview(client, actors["pm"], content)
    assert payload["authorities"]["ict"]["needs_look"] == 1
    assert payload["authorities"]["ict"]["in_tracker"] == 0


def test_one_unreadable_technology_makes_the_whole_village_need_a_look(client, actors):
    content = _fill(
        [(IDS["three_tech"], {"ict_2g": "yes", "ict_3g": "yes", "ict_4g": "??"})]
    )
    payload = _preview(client, actors["pm"], content)
    # Not "two thirds registered". Never an average.
    assert payload["authorities"]["ict"]["needs_look"] == 1


def test_a_village_is_in_the_tracker_only_when_every_requested_tech_is(client, actors):
    content = _fill(
        [(IDS["three_tech"], {"ict_2g": "yes", "ict_3g": "yes"})]  # 4G left blank
    )
    payload = _preview(client, actors["pm"], content)
    assert payload["authorities"]["ict"]["not_in_tracker"] == 1
    assert payload["authorities"]["ict"]["in_tracker"] == 0


def test_a_cell_for_a_technology_never_requested_changes_nothing(client, actors):
    """two_tech requested 3G and 4G. Its 2G column is not an answer.

    Filled with a yes it must not help; filled with nonsense it must not hurt.
    """
    for noise in ("yes", "???", "n/a"):
        content = _fill(
            [(IDS["two_tech"], {"ict_2g": noise, "ict_3g": "yes", "ict_4g": "yes"})]
        )
        payload = _preview(client, actors["pm"], content)
        assert payload["authorities"]["ict"]["in_tracker"] == 1, noise
        assert payload["authorities"]["ict"]["needs_look"] == 0, noise


def test_the_two_authorities_are_independent(client, actors):
    content = _fill(
        [
            (
                IDS["three_tech"],
                {
                    "ict_2g": "yes", "ict_3g": "yes", "ict_4g": "yes",
                    "cra_2g": "yes", "cra_3g": "", "cra_4g": "",
                },
            )
        ]
    )
    payload = _preview(client, actors["pm"], content)
    assert payload["authorities"]["ict"]["in_tracker"] == 1
    assert payload["authorities"]["cra"]["not_in_tracker"] == 1


def test_a_village_with_no_requested_technology_needs_a_look(client, actors):
    """We do not know what to look for, so we do not claim either answer."""
    content = _fill([(IDS["no_tech"], {"ict_2g": "yes"})])
    payload = _preview(client, actors["pm"], content)
    assert payload["authorities"]["ict"]["needs_look"] == 1


# ----- Rows that cannot be read ------------------------------------------


def test_unmatched_and_repeated_rows_are_listed_not_dropped(client, actors):
    content = _fill(
        [
            (IDS["three_tech"], {"ict_2g": "yes"}),
            (999_999, {"ict_2g": "yes"}),
            (IDS["three_tech"], {"ict_2g": "no"}),
            ("not-a-number", {"ict_2g": "yes"}),
        ]
    )
    payload = _preview(client, actors["pm"], content)
    assert payload["total_rows"] == 4
    assert payload["matched"] == 1
    assert payload["unmatched"] == 3
    reasons = {exception["reason"] for exception in payload["exceptions"]}
    assert "No village with that id" in reasons
    assert "village_id is not a number" in reasons
    assert any("Repeated village_id" in reason for reason in reasons)


def test_a_file_without_village_id_is_refused(client, actors):
    content = _fill(
        [(IDS["three_tech"], {})],
        headers=["site_id", "site_type", "ict_2g"],
    )
    response = _upload(client, actors["pm"], "import/preview", content)
    assert response.status_code == 400
    assert "village_id" in response.json()["detail"]


def test_something_that_is_not_a_workbook_is_refused(client, actors):
    response = _upload(client, actors["pm"], "import/preview", b"not a workbook")
    assert response.status_code == 400
    assert "template" in response.json()["detail"]


# ----- Nothing writes before Confirm -------------------------------------


def test_a_preview_writes_nothing_and_an_abandoned_upload_leaves_nothing(client, actors):
    from app.models.mojri import MojriImportRun, MojriTrackerStatus

    db = SessionLocal()
    before = (
        db.query(MojriTrackerStatus).count(),
        db.query(MojriImportRun).count(),
    )
    db.close()

    content = _fill(
        [(IDS["three_tech"], {"ict_2g": "yes", "ict_3g": "yes", "ict_4g": "yes"})]
    )
    payload = _preview(client, actors["pm"], content)
    assert payload["authorities"]["ict"]["moving_to_in_tracker"] == 1

    db = SessionLocal()
    after = (
        db.query(MojriTrackerStatus).count(),
        db.query(MojriImportRun).count(),
    )
    db.close()
    assert after == before, "the preview wrote to the database"
    assert _statuses(IDS["three_tech"]) is None


def test_confirm_refuses_a_file_that_is_not_the_one_previewed(client, actors):
    previewed = _fill([(IDS["three_tech"], {"ict_2g": "yes"})])
    payload = _preview(client, actors["pm"], previewed)

    swapped = _fill([(IDS["three_tech"], {"ict_2g": "yes", "ict_3g": "yes"})])
    response = _upload(
        client, actors["pm"], "import/commit", swapped, digest=payload["digest"]
    )
    assert response.status_code == 400
    assert "previewed" in response.json()["detail"]
    assert _statuses(IDS["three_tech"]) is None


# ----- Confirm ------------------------------------------------------------


def test_confirm_writes_exactly_what_the_preview_showed(client, actors):
    content = _fill(
        [
            (
                IDS["three_tech"],
                {
                    "ict_2g": "yes", "ict_3g": "yes", "ict_4g": "yes",
                    "cra_2g": "?", "cra_3g": "", "cra_4g": "",
                },
            ),
            (IDS["two_tech"], {"ict_3g": "yes", "ict_4g": ""}),
        ]
    )
    payload = _preview(client, actors["pm"], content)
    response = _upload(
        client, actors["pm"], "import/commit", content, digest=payload["digest"]
    )
    assert response.status_code == 200, response.text
    committed = response.json()

    assert committed["authorities"] == payload["authorities"]
    assert committed["matched"] == payload["matched"]
    assert _statuses(IDS["three_tech"]) == (IN_TRACKER, NEEDS_LOOK)
    assert _statuses(IDS["two_tech"]) == (NOT_IN_TRACKER, NOT_IN_TRACKER)

    # The run is recorded, and every status points at it.
    from app.models.mojri import MojriImportRun, MojriTrackerStatus

    db = SessionLocal()
    run = db.get(MojriImportRun, committed["import_run_id"])
    assert run is not None and run.matched_rows == 2
    assert {
        row.source_import_id
        for row in db.query(MojriTrackerStatus).all()
    } == {run.id}
    db.close()


def test_a_village_absent_from_this_file_is_flagged_and_never_reverted(client, actors):
    """three_tech is in the tracker from the import above. This file omits it."""
    assert _statuses(IDS["three_tech"])[0] == IN_TRACKER

    content = _fill([(IDS["two_tech"], {"ict_3g": "yes", "ict_4g": "yes"})])
    payload = _preview(client, actors["pm"], content)
    missing = {item["village_id"]: item for item in payload["disappeared"]}
    assert IDS["three_tech"] in missing
    assert missing[IDS["three_tech"]]["authorities"] == ["ICT"]

    response = _upload(
        client, actors["pm"], "import/commit", content, digest=payload["digest"]
    )
    assert response.status_code == 200, response.text
    # Still in the tracker. A row missing from a spreadsheet is not a
    # withdrawal, and this import says nothing about it.
    assert _statuses(IDS["three_tech"])[0] == IN_TRACKER
    assert _statuses(IDS["two_tech"]) == (IN_TRACKER, NOT_IN_TRACKER)


def test_the_import_is_a_full_snapshot_of_the_villages_it_names(client, actors):
    """two_tech was in the tracker; this file says it is not any more."""
    assert _statuses(IDS["two_tech"])[0] == IN_TRACKER

    content = _fill([(IDS["two_tech"], {"ict_3g": "", "ict_4g": ""})])
    payload = _preview(client, actors["pm"], content)
    response = _upload(
        client, actors["pm"], "import/commit", content, digest=payload["digest"]
    )
    assert response.status_code == 200, response.text
    assert _statuses(IDS["two_tech"])[0] == NOT_IN_TRACKER


def test_the_import_is_recorded_in_the_audit_log(client, actors):
    entries = client.get(
        "/api/v1/admin/audit-logs", headers=actors["admin"], params={"module": "mojri"}
    )
    assert entries.status_code == 200, entries.text
    rows = entries.json()
    rows = rows["items"] if isinstance(rows, dict) else rows
    assert any(row["action"] == "IMPORTED" for row in rows), rows


# ----- Who may do what ----------------------------------------------------


def test_admin_may_take_the_template_and_may_not_import(client, actors):
    """The Admin/PM separation, applied to this feature. Admin's button is a
    read; the write is PM's."""
    assert client.get(
        "/api/v1/mojri/template.xlsx", headers=actors["admin"]
    ).status_code == 200

    content = _fill([(IDS["three_tech"], {"ict_2g": "yes"})])
    assert _upload(
        client, actors["admin"], "import/preview", content
    ).status_code == 403
    assert _upload(
        client, actors["admin"], "import/commit", content, digest="x"
    ).status_code == 403


def test_pm_reaches_both(client, actors):
    assert client.get(
        "/api/v1/mojri/template.xlsx", headers=actors["pm"]
    ).status_code == 200
    content = _fill([(IDS["three_tech"], {"ict_2g": "yes"})])
    assert _upload(client, actors["pm"], "import/preview", content).status_code == 200


def test_coordinator_reaches_neither(client, actors):
    assert client.get(
        "/api/v1/mojri/template.xlsx", headers=actors["coordinator"]
    ).status_code == 403
    content = _fill([(IDS["three_tech"], {"ict_2g": "yes"})])
    assert _upload(
        client, actors["coordinator"], "import/preview", content
    ).status_code == 403


# ----- What this feature must not touch -----------------------------------


def test_acceptance_is_untouched_by_an_import(client, actors):
    """This is a parallel record of somebody else's paperwork. It reads our
    acceptance data and never writes it."""
    from app.models.workitem import Village

    db = SessionLocal()
    before = {
        village.id: (village.ict_status, village.cra_status)
        for village in db.query(Village).all()
    }
    db.close()

    content = _fill(
        [(IDS["cra_only"], {"cra_2g": "yes", "ict_2g": "yes"})]
    )
    payload = _preview(client, actors["pm"], content)
    _upload(client, actors["pm"], "import/commit", content, digest=payload["digest"])

    db = SessionLocal()
    after = {
        village.id: (village.ict_status, village.cra_status)
        for village in db.query(Village).all()
    }
    db.close()
    assert after == before


def test_mojri_is_not_a_third_acceptance_authority():
    """The value set the acceptance workflow knows about is unchanged."""
    from app.models.acceptance_workflow import AUTHORITIES

    assert AUTHORITIES == ("ICT", "CRA")
