"""Lifecycle Gaps export (``GET /gaps/villages.xlsx``): the villages behind a
figure, as a spreadsheet.

The property that matters: **the file for a figure has that figure's row
count.** It is asserted for every exportable gap under every lens, against
every owner row the overview returns, and against the country total -- so a
drawer row, the tile above it and the file behind either cannot disagree.

The data is the overview's hand-worked seed (``test_gaps_overview._seed``),
which already holds every awkward case: a province-less site, a work item
with no contractor, a repeated village code, and villages the universe must
exclude.

Run with:  cd backend && pytest tests/test_gaps_export.py -q
"""
# The `actors` fixture is imported from the overview tests and then named as a
# test parameter, which is how pytest injects it; pyflakes reads that as a
# redefinition.
# ruff: noqa: F811
import io
import os
import sys

import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_gaps_export_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.services import gap_export, gaps  # noqa: E402
from tests.conftest import create_schema  # noqa: E402
from tests.test_gaps_overview import (  # noqa: E402, F401  (actors is a fixture)
    ALPHA,
    GAPS,
    MAZANDARAN,
    TEHRAN,
    TOTALS,
    _overview,
    _seed,
    actors,
)

URL = "/api/v1/gaps/villages.xlsx"


@pytest.fixture(scope="module")
def client():
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        _seed()
        yield c


def _export(client, headers, **params):
    response = client.get(URL, headers=headers, params=params)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    return response, book


def _villages(book) -> list[dict]:
    rows = list(book["Villages"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    return [dict(zip(header, row, strict=True)) for row in body]


def _count(client, headers, **params) -> int:
    response, book = _export(client, headers, **params)
    count = len(_villages(book))
    assert int(response.headers["x-village-count"]) == count
    return count


# ----- One figure, one file ----------------------------------------------


def _figure(payload: dict, gap: str) -> int:
    if gap in payload["gaps"]:
        return payload["gaps"][gap]["count"]
    return payload["totals"][gap]


@pytest.mark.parametrize("gap", list(gaps.EXPORTS))
def test_the_country_file_has_the_country_figure(client, actors, gap):
    payload = _overview(client, actors["pm"])
    assert _count(client, actors["pm"], gap=gap) == _figure(payload, gap)


@pytest.mark.parametrize("lens", gaps.LENSES)
def test_every_owner_row_exports_exactly_its_count(client, actors, lens):
    """Every gap × every row of this lens: the drawer's number and the file
    behind it agree, attribution rows included."""
    payload = _overview(client, actors["pm"], lens=lens)
    problems = []
    for gap in gaps.GAP_KEYS:
        for row in payload["rows"][gap]:
            got = _count(client, actors["pm"], gap=gap, lens=lens, key=row["name"])
            if got != row["count"]:
                problems.append(f"{lens}/{gap}/{row['name']}: file {got}, row {row['count']}")
    assert not problems, "\n".join(problems)


def test_the_hand_worked_figures_come_out_of_the_files(client, actors):
    for gap, figure in GAPS.items():
        assert _count(client, actors["pm"], gap=gap) == figure["count"], gap
    assert _count(client, actors["pm"], gap="ict_approved") == TOTALS["ict_approved"]
    assert _count(client, actors["pm"], gap="cra_approved") == TOTALS["cra_approved"]


def test_attribution_rows_are_exported_not_dropped(client, actors):
    _, book = _export(client, actors["pm"], gap="pending_ict", lens="province", key="Unknown province")
    assert book["Villages"].max_row == 1  # WI3 is fully approved: no pending ICT
    _, book = _export(client, actors["pm"], gap="ict_approved", lens="province", key="Unknown province")
    rows = _villages(book)
    assert len(rows) == 2
    assert {row["Attribution"] for row in rows} == {gaps.UNKNOWN_PROVINCE}
    assert {row["Province"] for row in rows} == {"Unknown province"}

    _, book = _export(client, actors["pm"], gap="pending_ict", lens="contractor", key="Unassigned")
    rows = _villages(book)
    assert len(rows) == 1
    assert rows[0]["Contractor"] == "Unassigned"
    assert rows[0]["Attribution"] == gaps.UNASSIGNED


def test_a_repeated_village_code_is_two_rows(client, actors):
    _, book = _export(client, actors["pm"], gap="ict_approved")
    codes = [row["Village ID"] for row in _villages(book)]
    assert codes.count("V1-00") == 2


# ----- Map scopes ---------------------------------------------------------


def test_a_province_scope_is_that_provinces_villages(client, actors):
    # Pending ICT by province, hand-worked: Tehran 9, Mazandaran 1.
    assert _count(client, actors["pm"], gap="pending_ict", scope=f"province:{TEHRAN}") == 9
    assert _count(client, actors["pm"], gap="pending_ict", scope=f"province:{MAZANDARAN}") == 1


def test_a_region_scope_is_the_sum_of_its_provinces(client, actors):
    # Tehran and Mazandaran are both North.
    assert _count(client, actors["pm"], gap="pending_ict", scope="region:North") == 10


def test_a_scope_and_an_owner_narrow_together(client, actors):
    assert _count(
        client, actors["pm"], gap="pending_ict", scope=f"province:{TEHRAN}",
        lens="contractor", key=ALPHA,
    ) == 8


# ----- The workbook -------------------------------------------------------


def test_the_villages_sheet_has_the_columns_a_filter_and_a_frozen_header(client, actors):
    _, book = _export(client, actors["pm"], gap="pending_cra")
    assert book.sheetnames == ["Villages", "Summary"]
    sheet = book["Villages"]
    header = [cell.value for cell in sheet[1]]
    assert header == [title for title, _, _ in gap_export.COLUMNS]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:M{GAPS['pending_cra']['count'] + 1}"


def test_the_row_says_who_owns_the_village_and_where_it_stands(client, actors):
    _, book = _export(client, actors["pm"], gap="cra_remained", lens="province", key="Mazandaran")
    [row] = _villages(book)
    assert row["Village ID"] == "V2-0"
    assert row["Province"] == "Mazandaran"
    assert row["CRA region"] == "North"
    assert row["Regional manager"] == "Nobakht"
    assert row["Coordinator"] == "Amir"
    assert row["Contractor"] == "DT-Beta"
    assert (row["ICT status"], row["CRA status"]) == ("Approved", "Pending")
    assert row["Mojri status"] == "ICT missing · CRA missing"
    assert row["Attribution"] == gaps.OWNED


def test_the_summary_says_what_the_file_is(client, actors):
    _, book = _export(client, actors["pm"], gap="pending_ict", lens="coordinator", key="Amir")
    summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True)}
    assert summary["Figure"] == "Pending ICT approval"
    assert summary["Filter"] == "PSO Coordinator = Amir"
    assert summary["Villages"] == 10
    assert summary["Exported (Jalali)"].startswith("14")
    assert summary["Exported (Gregorian)"].endswith("UTC")
    assert "test_pm" in summary["Exported by"]

    _, book = _export(client, actors["pm"], gap="pending_ict")
    summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True)}
    assert summary["Filter"] == "All villages"


def test_the_filename_is_ascii_and_names_the_filter(client, actors):
    response, _ = _export(client, actors["pm"], gap="pending_cra", lens="coordinator", key="Amir")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="uep-pending-cra-coordinator-amir-14')
    assert disposition.isascii()

    response, _ = _export(client, actors["pm"], gap="pending_ict", scope=f"province:{TEHRAN}")
    assert 'filename="uep-pending-ict-province-tehran-' in response.headers["content-disposition"]

    response, _ = _export(client, actors["pm"], gap="pending_ict", lens="contractor", key="پیشرو فن")
    assert response.headers["content-disposition"].isascii()


def test_a_large_export_is_streamed_and_still_whole(client, actors, monkeypatch):
    monkeypatch.setattr(gap_export, "STREAM_THRESHOLD_ROWS", 1)
    response, book = _export(client, actors["pm"], gap="pending_ict")
    assert "content-length" not in response.headers
    assert len(_villages(book)) == GAPS["pending_ict"]["count"]


# ----- Who may export what ------------------------------------------------


def test_admin_is_refused(client, actors):
    assert client.get(URL, headers=actors["admin"], params={"gap": "pending_ict"}).status_code == 403


def test_a_non_pm_gets_only_their_own_villages(client, actors):
    # RM Allahyar owns Tehran: 9 pending ICT (see the overview test).
    assert _count(client, actors["rm"], gap="pending_ict") == 9
    assert _count(client, actors["rm"], gap="pending_ict", lens="rm", key="Allahyar") == 9
    # Contractor Alpha: WI1 + WI3.
    assert _count(client, actors["contractor"], gap="pending_ict") == 8


@pytest.mark.parametrize(
    "actor, params",
    [
        ("rm", {"lens": "rm", "key": "Pirayesh"}),
        ("rm", {"lens": "coordinator", "key": "Hossein"}),
        ("rm", {"lens": "province", "key": "Mazandaran"}),
        ("coordinator", {"lens": "coordinator", "key": "Hossein"}),
        ("contractor", {"lens": "contractor", "key": "DT-Beta"}),
    ],
)
def test_a_non_pm_asking_for_somebody_elses_villages_is_refused(client, actors, actor, params):
    response = client.get(URL, headers=actors[actor], params={"gap": "pending_ict", **params})
    assert response.status_code == 403, response.text


def test_a_non_pm_may_narrow_within_their_own_villages(client, actors):
    # Coordinator Amir holds Tehran and Mazandaran; Mazandaran is his.
    assert _count(
        client, actors["coordinator"], gap="pending_ict", lens="province", key="Mazandaran"
    ) == 1


@pytest.mark.parametrize(
    "params",
    [
        {"gap": "everything"},
        {"gap": "pending_ict", "lens": "country", "key": "x"},
        {"gap": "pending_ict", "lens": "rm"},
        {"gap": "pending_ict", "key": "Allahyar"},
        {"gap": "pending_ict", "scope": "city:Tehran"},
        {"gap": "pending_ict", "scope": "province:"},
    ],
)
def test_a_malformed_request_is_refused(client, actors, params):
    assert client.get(URL, headers=actors["pm"], params=params).status_code == 422
