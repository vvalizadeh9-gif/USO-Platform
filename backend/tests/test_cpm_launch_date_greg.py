"""The Gregorian launch-date column, through the real import pipeline.

The DT dashboard's "Where this is going" chart read on-aired as zero because
``launch_date_gregorian`` was never populated by the import (fixed in
cpm_import._master_fields) and because ``CpmImportService._date`` did not
recognise every shape ``pandas.read_excel`` can hand back for a real date
cell. This exercises the actual pandas round-trip -- not a Python value
handed straight to the ORM -- so a regression in either place is caught here
rather than in production months later.

Run with:  cd backend && pytest tests/test_cpm_launch_date_greg.py -q
"""
import os
import sys
from datetime import date

import openpyxl
import pytest

os.environ["DATABASE_URL"] = "sqlite:////tmp/uep_cpmgreg_pytest.db"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import sqlalchemy.dialects.postgresql as _pg  # noqa: E402
from sqlalchemy import JSON  # noqa: E402

_pg.JSONB = JSON

from fastapi.testclient import TestClient  # noqa: E402

from app.core import jalali  # noqa: E402
from app.core.database import SessionLocal  # noqa: E402
from tests.conftest import create_schema, sample_cpm_path  # noqa: E402

SRC = sample_cpm_path()


@pytest.fixture(scope="module")
def client():
    if os.path.exists("/tmp/uep_cpmgreg_pytest.db"):
        os.remove("/tmp/uep_cpmgreg_pytest.db")
    create_schema()
    from app.main import app

    with TestClient(app) as c:
        yield c


def test_real_date_cell_becomes_launch_date_gregorian(client, tmp_path):
    """A genuine Excel date cell -- written and formatted the way Excel
    itself stores one, then read back through pandas -- must survive the
    import into WorkItem.launch_date_gregorian, and onair_month() must place
    it in the past rather than dropping it."""
    if not SRC:
        pytest.skip("no sample CPM workbook available")
    from app.models.reference import User
    from app.models.workitem import Site, WorkItem
    from app.services import cpm_columns as C
    from app.services.cpm_import import CpmImportService
    from app.services.drive_test_analytics import onair_month

    wb = openpyxl.load_workbook(SRC)
    ws = wb["CPM"]
    target_excel_row = 4  # header at index 2 (0-based) => data starts row 4
    ws.cell(
        row=target_excel_row, column=C.COL["last_stage"] + 1, value="راه_اندازی_دائم"
    )
    date_cell = ws.cell(row=target_excel_row, column=C.COL["launch_date_greg"] + 1)
    date_cell.value = date(2021, 11, 3)
    date_cell.number_format = "mm/dd/yyyy"
    site_code = ws.cell(row=target_excel_row, column=C.COL["site_code"] + 1).value

    mod = str(tmp_path / "greg_date.xlsx")
    wb.save(mod)

    db = SessionLocal()
    try:
        admin = db.query(User).filter(User.username == "admin").first()
        CpmImportService(db, user_id=admin.id).import_file(mod, "greg_date.xlsx")

        site = db.query(Site).filter(Site.site_code == site_code).first()
        wi = db.query(WorkItem).filter(WorkItem.site_id == site.id).first()

        assert wi.launch_date_gregorian == date(2021, 11, 3)
        assert onair_month(wi) == jalali.to_shamsi(date(2021, 11, 3))
    finally:
        db.close()


def test_bare_excel_serial_number_also_parses() -> None:
    """The exact shape the bug report's CPM file was actually in: the cell's
    number format was never set to a date, so pandas hands back a plain
    day-count rather than a Timestamp. This must not read as blank."""
    from app.services.cpm_import import CpmImportService

    # 44503 is the Excel serial for 2021-11-03 (days since 1899-12-30).
    assert CpmImportService._date(44503) == date(2021, 11, 3)
    assert CpmImportService._date(44503.0) == date(2021, 11, 3)


def test_header_mismatch_is_detected() -> None:
    from app.services import cpm_columns as C

    good_headers = [""] * 30
    good_headers[C.COL["site_code"]] = "کد سایت ایرانسل"
    good_headers[C.COL["site_type"]] = "نوع سایت"
    good_headers[C.COL["last_stage"]] = "آخرین مرحله انجام شده"
    good_headers[C.COL["launch_date_shamsi"]] = "تاریخ راه اندازی"
    good_headers[C.COL["launch_date_greg"]] = "تاریخ راه اندازی (میلادی)"
    good_headers[C.COL["project_name"]] = "نام پروژه اجرایی سایت"
    good_headers[C.COL["pm_name"]] = "مدیر پروژه"
    assert C.find_header_mismatches(good_headers) == []

    shifted = list(good_headers)
    shifted.insert(C.COL["launch_date_shamsi"], "یک ستون جدید")
    mismatches = C.find_header_mismatches(shifted)
    assert any("launch_date_greg" in m for m in mismatches)
